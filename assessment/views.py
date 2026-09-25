from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.core.exceptions import PermissionDenied
from django.db import transaction
from django.http import HttpResponseRedirect, JsonResponse
from django.urls import reverse
from django.utils import timezone
from django.db.models import Prefetch, Count, Sum, Max, Q

from accounts.decorators import role_required
from accounts.models import Organization
from .models import (
    AwardCycle, QuestionnaireTemplate, AssessmentCategory, Criterion, LevelIndicator,
    Questionnaire, Response, VerifierResponse, JudgeResponse,
    EvidenceDocument, EvidenceLink, StageSubmission,
)
from .forms import CycleForm, BulkResponseForm, VerifierScoreForm, JudgeExcelForm
from .excel_parser import parse_excel_questionnaire
from .coverage import criterion_coverage, category_coverage, questionnaire_coverage
from audit.models import AuditLog

import os as _os
import re as _re
import uuid as _uuid

_ALLOWED_EXCEL_EXT = {'.xlsx', '.xls'}
_ALLOWED_EVIDENCE_EXT = {'.pdf', '.doc', '.docx', '.xls', '.xlsx', '.png', '.jpg', '.jpeg'}
_MAX_EXCEL_BYTES = 20 * 1024 * 1024   # 20 MB


def _categories_for_questionnaire(questionnaire, *, active_only=True):
    """Return only categories assigned to this questionnaire's template.

    A questionnaire without a template must not fall back to cycle-wide
    categories, because that can expose another organization's assessment.
    """
    if not questionnaire.template_id:
        return AssessmentCategory.objects.none()

    filters = {
        'cycle_id': questionnaire.cycle_id,
        'template_id': questionnaire.template_id,
        'is_informal_sector_only': (
            questionnaire.organization.org_type == 'informal_sector'
        ),
    }
    if active_only:
        filters['is_active'] = True

    return AssessmentCategory.objects.filter(**filters).order_by('order', 'name')

# Magic-byte signatures keyed by extension. Upload content must start with one of
# these for the claimed extension, so a renamed file (e.g. evil.html -> evil.png) is rejected.
_ZIP_SIGS = [b'PK\x03\x04']
_OLE_SIG = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'
_FILE_SIGNATURES = {
    '.pdf':  [b'%PDF'],
    '.png':  [b'\x89PNG\r\n\x1a\n'],
    '.jpg':  [b'\xff\xd8\xff'],
    '.jpeg': [b'\xff\xd8\xff'],
    '.docx': _ZIP_SIGS,
    '.xlsx': _ZIP_SIGS,
    '.doc':  [_OLE_SIG],
    '.xls':  [_OLE_SIG],
}


def _content_matches_extension(f):
    """True if the uploaded file's leading bytes match its claimed extension.
    Restores the stream position so the file can still be saved afterwards."""
    ext = _os.path.splitext(f.name)[1].lower()
    sigs = _FILE_SIGNATURES.get(ext)
    if not sigs:
        return False
    try:
        pos = f.tell()
    except (AttributeError, OSError):
        pos = 0
    header = b''
    try:
        f.seek(0)
        header = f.read(8)
    finally:
        try:
            f.seek(pos)
        except (AttributeError, OSError):
            pass
    return any(header.startswith(s) for s in sigs)


def _validate_excel_upload(f):
    """Return an error string or None if the file is acceptable."""
    ext = _os.path.splitext(f.name)[1].lower()
    if ext not in _ALLOWED_EXCEL_EXT:
        return f'"{f.name}" is not a supported Excel file (.xlsx or .xls).'
    if f.size > _MAX_EXCEL_BYTES:
        return f'"{f.name}" exceeds the 20 MB size limit.'
    if not _content_matches_extension(f):
        return f'"{f.name}" — file content does not match its extension.'
    return None



def _validate_judge_excel_upload(f):
    """Validate a judge Excel attachment."""
    if not f:
        return None
    ext = _os.path.splitext(f.name)[1].lower()
    if ext not in _ALLOWED_EXCEL_EXT:
        return f'"{f.name}" is not a supported Excel file (.xlsx or .xls).'
    if f.size > _MAX_EXCEL_BYTES:
        return f'"{f.name}" exceeds the 20 MB size limit.'
    if not _content_matches_extension(f):
        return f'"{f.name}" — file content does not match its extension.'
    return None


def _validate_evidence_upload(f, questionnaire, response):
    """Return an error string or None. Org-level quota + per-file + per-criterion."""
    from django.conf import settings
    from django.template.defaultfilters import filesizeformat
    ext = _os.path.splitext(f.name)[1].lower()
    if ext not in _ALLOWED_EVIDENCE_EXT:
        return f'"{f.name}" — unsupported file type.'
    if not _content_matches_extension(f):
        return f'"{f.name}" — file content does not match its extension.'
    if f.size > settings.EVIDENCE_MAX_FILE_BYTES:
        return f'"{f.name}" — exceeds {filesizeformat(settings.EVIDENCE_MAX_FILE_BYTES)} limit.'
    if response.evidence_links.count() >= settings.EVIDENCE_MAX_FILES_PER_RESP:
        return f'"{f.name}" — max {settings.EVIDENCE_MAX_FILES_PER_RESP} files per criterion.'
    used = (
        EvidenceDocument.objects
        .filter(organization=questionnaire.organization)
        .aggregate(t=Sum('file_size'))['t'] or 0
    )
    if used + f.size > settings.EVIDENCE_MAX_QUOTA_BYTES:
        return f'"{f.name}" — library quota ({filesizeformat(settings.EVIDENCE_MAX_QUOTA_BYTES)}) would be exceeded.'
    return None


# ── Admin ────────────────────────────────────────────────────────────────────

@role_required('admin')
def admin_dashboard(request):
    cycles = AwardCycle.objects.all()
    open_cycle = cycles.filter(is_open=True).first()
    total_orgs = Organization.objects.filter(is_active=True).count()
    submitted = 0
    total_q = 0
    org_data = []
    submission_rate = 0
    if open_cycle:
        qs = open_cycle.questionnaires.select_related('organization').all()
        total_q = qs.count()
        submitted = qs.filter(is_submitted=True).count()
        submission_rate = round(submitted / total_q * 100) if total_q else 0

        verified_counts = (
            JudgeResponse.objects
            .filter(questionnaire__cycle=open_cycle, score__isnull=False)
            .values('questionnaire_id')
            .annotate(count=Count('id'))
        )
        verified_map = {v['questionnaire_id']: v['count'] for v in verified_counts}

        answered_counts = (
            Response.objects
            .filter(questionnaire__cycle=open_cycle, score__isnull=False)
            .values('questionnaire_id')
            .annotate(count=Count('id'))
        )
        answered_map = {r['questionnaire_id']: r['count'] for r in answered_counts}

        for q in qs:
            org_data.append({
                'org': q.organization,
                'questionnaire': q,
                'answered_count': answered_map.get(q.pk, 0),
                'verified_count': verified_map.get(q.pk, 0),
            })
    return render(request, 'assessment/dashboard_admin.html', {
        'open_cycle': open_cycle,
        'total_orgs': total_orgs,
        'submitted': submitted,
        'total_q': total_q,
        'cycles': cycles,
        'org_data': org_data,
        'submission_rate': submission_rate,
    })


@role_required('admin')
def cycle_list(request):
    cycles = AwardCycle.objects.all()
    return render(request, 'assessment/cycle_list.html', {'cycles': cycles})


@role_required('admin')
def cycle_create(request):
    form = CycleForm(request.POST or None, request.FILES or None)
    if request.method == 'POST' and form.is_valid():
        cycle = form.save()
        if cycle.excel_file:
            err = _validate_excel_upload(cycle.excel_file)
            if err:
                messages.error(request, err)
                return redirect('cycle_detail', pk=cycle.pk)
            warnings = parse_excel_questionnaire(cycle.excel_file, cycle)
            for w in warnings:
                messages.warning(request, w)
            messages.success(request, f'Cycle "{cycle.name}" created and questionnaire imported.')
        else:
            messages.success(request, f'Cycle "{cycle.name}" created.')
        return redirect('cycle_detail', pk=cycle.pk)
    return render(request, 'assessment/cycle_form.html', {'form': form, 'action': 'Create'})


@role_required('admin')
def cycle_detail(request, pk):
    cycle = get_object_or_404(AwardCycle, pk=pk)
    if request.method == 'POST' and 'excel_file' in request.FILES:
        if cycle.is_open:
            messages.error(request, 'Cannot re-upload questionnaire after the cycle has been opened.')
            return redirect('cycle_detail', pk=cycle.pk)
        upload = request.FILES['excel_file']
        err = _validate_excel_upload(upload)
        if err:
            messages.error(request, err)
            return redirect('cycle_detail', pk=cycle.pk)
        cycle.excel_file = upload
        cycle.save(update_fields=['excel_file'])
        warnings = parse_excel_questionnaire(cycle.excel_file, cycle)
        for w in warnings:
            messages.warning(request, w)
        messages.success(request, 'Questionnaire re-imported.')
        return redirect('cycle_detail', pk=cycle.pk)
    questionnaires = cycle.questionnaires.select_related('organization').order_by('organization__name')
    categories = list(
        cycle.categories.filter(is_active=True).prefetch_related(
            Prefetch(
                'criteria',
                queryset=Criterion.objects.filter(is_active=True).order_by('number'),
                to_attr='active_criteria',
            )
        )
    )
    assigned_org_ids = questionnaires.values_list('organization_id', flat=True)
    missing_orgs = Organization.objects.filter(is_active=True).exclude(id__in=assigned_org_ids)
    return render(request, 'assessment/cycle_detail.html', {
        'cycle': cycle,
        'questionnaires': questionnaires,
        'categories': categories,
        'missing_orgs': missing_orgs,
        'templates': cycle.templates.filter(is_active=True).order_by('name'),
    })


@role_required('admin')
def cycle_toggle_open(request, pk):
    cycle = get_object_or_404(AwardCycle, pk=pk)
    if request.method == 'POST':
        cycle.is_open = not cycle.is_open
        cycle.save()
        state = 'opened' if cycle.is_open else 'closed'
        action = 'cycle.opened' if cycle.is_open else 'cycle.closed'
        AuditLog.log(request, action,
            f'{state.capitalize()} award cycle "{cycle.name}"',
            object_type='AwardCycle', object_id=cycle.pk, object_repr=cycle.name)
        messages.success(request, f'Cycle "{cycle.name}" {state}.')
    return redirect('cycle_detail', pk=cycle.pk)


@role_required('admin')
def cycle_add_org_questionnaire(request, pk):
    cycle = get_object_or_404(AwardCycle, pk=pk, is_open=True)
    if request.method == 'POST':
        org_id = request.POST.get('org_id', '')
        if not org_id or not str(org_id).isdigit():
            messages.error(request, 'Invalid organization selected.')
            return redirect('cycle_detail', pk=pk)
        org = get_object_or_404(Organization, pk=int(org_id), is_active=True)
        template_id = request.POST.get('template_id', '').strip()
        template = None
        if template_id:
            template = get_object_or_404(QuestionnaireTemplate, pk=template_id, cycle=cycle, is_active=True)
        elif org.assessment_profile:
            template = QuestionnaireTemplate.objects.filter(cycle=cycle, assessment_profile=org.assessment_profile, is_active=True).first()
        if not template:
            messages.error(request, 'Select a questionnaire template for this organization before assigning it.')
            return redirect('cycle_detail', pk=pk)
        _, created = Questionnaire.objects.get_or_create(
            cycle=cycle, organization=org, defaults={'template': template}
        )
        if not created and not _.template_id:
            _.template = template
            _.save(update_fields=['template', 'updated_at'])
        if created:
            messages.success(request, f'Questionnaire created for {org.name}.')
        else:
            messages.info(request, f'{org.name} already has a questionnaire for this cycle.')
    return redirect('cycle_detail', pk=pk)


@role_required('admin')
def category_edit(request, pk):
    category = get_object_or_404(AssessmentCategory, pk=pk)
    if request.method == 'POST':
        from decimal import Decimal, InvalidOperation
        name = request.POST.get('name', '').strip()
        code = request.POST.get('code', '').strip()
        raw_weight = request.POST.get('weight', '').strip()
        if not name or not code:
            messages.error(request, 'Category code and name cannot be empty.')
            return redirect('category_edit', pk=category.pk)
        try:
            weight = Decimal(raw_weight or '0')
            if weight < 0:
                raise InvalidOperation
        except (InvalidOperation, ValueError):
            messages.error(request, 'Enter a valid non-negative category weight.')
            return redirect('category_edit', pk=category.pk)
        category.code = code
        category.name = name
        category.weight = weight
        category.save(update_fields=['code', 'name', 'weight', 'updated_at'])
        messages.success(request, f'Category {code} updated.')
        return redirect('cycle_detail', pk=category.cycle_id)
    return render(request, 'assessment/category_edit.html', {'category': category})


@role_required('admin')
def cycle_results(request, pk):
    from .analytics import compute_cycle_analytics
    cycle = get_object_or_404(AwardCycle, pk=pk)
    analytics = compute_cycle_analytics(cycle)
    return render(request, 'assessment/cycle_results.html', {
        'cycle': cycle,
        'analytics': analytics,
    })


@login_required
def org_report(request, pk, q_pk):
    from accounts.models import VerifierAssignment
    from .analytics import compute_org_analytics, compute_org_benchmark
    cycle = get_object_or_404(AwardCycle, pk=pk)
    questionnaire = get_object_or_404(Questionnaire, pk=q_pk, cycle=cycle, is_submitted=True)
    profile = request.user.userprofile
    role = profile.role

    if role == 'admin':
        pass
    elif role == 'verifier':
        is_assigned = VerifierAssignment.objects.filter(
            verifier=request.user, organization=questionnaire.organization
        ).exists()
        if not is_assigned or cycle.is_open:
            raise PermissionDenied
    elif role == 'member':
        if (
            profile.organization != questionnaire.organization or
            cycle.is_open or
            not questionnaire.is_distributed
        ):
            raise PermissionDenied
    else:
        raise PermissionDenied

    # Cohort context carries this org's rank and the league averages only —
    # compute_org_benchmark never returns another organisation's name or score,
    # which is what makes it safe to show a member.
    benchmark = compute_org_benchmark(questionnaire)
    analytics = compute_org_analytics(questionnaire, benchmark=benchmark)
    return render(request, 'assessment/org_report.html', {
        'cycle': cycle,
        'questionnaire': questionnaire,
        'analytics': analytics,
        'benchmark': benchmark,
        'is_admin': role == 'admin',
        'is_member': role == 'member',
    })


_ORG_SHEET_HEADERS = [
    'Category', 'Criterion #', 'Criterion Name',
    'Verified Score', 'Verified Level', 'Self-Assessed', 'Self-Awareness',
    'Why this level', 'What the next level requires', 'Verifier comments',
]


def _format_verifier_notes(notes):
    """One line per verifier, named, with the score that comment justifies."""
    lines = []
    for v in notes:
        who = v['name']
        if v['score'] is not None:
            who = f"{who} (scored {v['score']})"
        lines.append(f"{who}: {v['notes']}")
    return '\n'.join(lines)


def _write_org_sheet(ws, Font, cycle, questionnaire, analytics):
    """One organisation's full detail, including the reasons behind each score.

    The verifier's written rationale and the level indicators are the whole
    point of the report, so the export carries them rather than reducing the
    assessment back down to three bare numbers.
    """
    ws.append(['Organization', questionnaire.organization.name])
    ws.append(['Cycle', cycle.name])
    ws.append(['Verified Score (out of 5)', analytics['overall_verified_score']])
    ws.append(['Self-Assessed Score', analytics['overall_claimed_score']])
    ws.append(['Self-Awareness', analytics['gap_phrase']])
    ws.append([
        'Questionnaire Completed',
        f"{analytics['completeness']}% "
        f"({analytics['answered_count']} of {analytics['criteria_count']} criteria)",
    ])
    ws.append([])
    ws.append(_ORG_SHEET_HEADERS)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True)

    for cat_entry in analytics['categories']:
        for e in cat_entry['criteria']:
            ws.append([
                cat_entry['category'].name,
                e['criterion'].number,
                e['criterion'].name,
                e['verified_score'],
                e['level_now'],
                e['claimed_score'],
                e['gap_phrase'],
                e['level_now_text'] or '',
                e['level_next_text'] or '',
                _format_verifier_notes(e['verifier_notes']),
            ])
    return ws


@role_required('admin')
def export_cycle_excel(request, pk):
    import openpyxl
    from openpyxl.styles import Font
    from django.http import HttpResponse
    from .analytics import compute_cycle_analytics, compute_org_analytics

    cycle = get_object_or_404(AwardCycle, pk=pk)
    analytics = compute_cycle_analytics(cycle)

    wb = openpyxl.Workbook()
    ws_rank = wb.active
    ws_rank.title = 'Rankings'
    ws_rank.append([
        'League', 'Rank', 'Organization', 'Verified Score', 'Self-Assessed',
        'Self-Awareness', 'Completed %', 'Criteria Answered',
    ])
    for cell in ws_rank[1]:
        cell.font = Font(bold=True)
    for league in analytics['leagues']:
        for entry in league['ranked_orgs']:
            ws_rank.append([
                league['label'],
                entry['rank'],
                entry['org'].name,
                entry['overall_verified_score'],
                entry['overall_claimed_score'],
                entry['gap_phrase'],
                entry['completeness'],
                f"{entry['answered']} of {entry['total']}",
            ])

    # Per-category standings: every organisation's position in every category,
    # which the old best/worst-only sheet could not express.
    ws_cats = wb.create_sheet(title='By Category')
    ws_cats.append(['League', 'Category', 'Rank', 'Organization', 'Verified Score'])
    for cell in ws_cats[1]:
        cell.font = Font(bold=True)
    for entry in analytics['category_table']:
        for place in entry['placings']:
            ws_cats.append([
                entry['league_label'],
                entry['category'].name,
                place['rank'],
                place['org'].name,
                place['score'],
            ])

    used_sheet_names = {}
    for entry in analytics['ranked_orgs']:
        q = entry['questionnaire']
        org_analytics = compute_org_analytics(q)
        base = q.organization.name[:31]
        count = used_sheet_names.get(base, 0)
        sheet_name = base if count == 0 else f"{base[:28]}-{count}"
        used_sheet_names[base] = count + 1
        ws_org = wb.create_sheet(title=sheet_name)
        _write_org_sheet(ws_org, Font, cycle, q, org_analytics)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="cycle-{cycle.year}-results.xlsx"'
    wb.save(response)
    return response


@role_required('admin', 'verifier')
def export_org_excel(request, pk, q_pk):
    import re
    import openpyxl
    from openpyxl.styles import Font
    from django.http import HttpResponse
    from accounts.models import VerifierAssignment
    from .analytics import compute_org_analytics

    cycle = get_object_or_404(AwardCycle, pk=pk)
    questionnaire = get_object_or_404(Questionnaire, pk=q_pk, cycle=cycle, is_submitted=True)
    profile = request.user.userprofile
    is_admin = profile.role == 'admin'
    if profile.role == 'verifier' and cycle.is_open:
        raise PermissionDenied
    is_assigned = VerifierAssignment.objects.filter(
        verifier=request.user, organization=questionnaire.organization
    ).exists()
    if not (is_admin or is_assigned):
        raise PermissionDenied

    analytics = compute_org_analytics(questionnaire)
    org_name = questionnaire.organization.name

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = org_name[:31]
    _write_org_sheet(ws, Font, cycle, questionnaire, analytics)

    slug = re.sub(r'[^\w\-]', '-', org_name.lower())[:30]
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{slug}-{cycle.year}-report.xlsx"'
    wb.save(response)
    return response


@role_required('admin', 'verifier')
def export_org_pdf(request, pk, q_pk):
    import re
    from io import BytesIO
    from django.http import HttpResponse
    from django.template.loader import render_to_string
    from xhtml2pdf import pisa
    from accounts.models import VerifierAssignment
    from .analytics import compute_org_analytics, compute_org_benchmark

    cycle = get_object_or_404(AwardCycle, pk=pk)
    questionnaire = get_object_or_404(Questionnaire, pk=q_pk, cycle=cycle, is_submitted=True)
    profile = request.user.userprofile
    is_admin = profile.role == 'admin'
    if profile.role == 'verifier' and cycle.is_open:
        raise PermissionDenied
    is_assigned = VerifierAssignment.objects.filter(
        verifier=request.user, organization=questionnaire.organization
    ).exists()
    if not (is_admin or is_assigned):
        raise PermissionDenied

    benchmark = compute_org_benchmark(questionnaire)
    analytics = compute_org_analytics(questionnaire, benchmark=benchmark)
    html = render_to_string('assessment/org_report_pdf.html', {
        'cycle': cycle,
        'questionnaire': questionnaire,
        'analytics': analytics,
        'benchmark': benchmark,
    }, request=request)

    pdf_buffer = BytesIO()
    result = pisa.CreatePDF(html, dest=pdf_buffer)
    if result.err:
        from django.http import HttpResponseServerError
        return HttpResponseServerError('PDF generation failed.')

    org_name = questionnaire.organization.name
    slug = re.sub(r'[^\w\-]', '-', org_name.lower())[:30]
    response = HttpResponse(pdf_buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{slug}-{cycle.year}-report.pdf"'
    return response


# ── Member ───────────────────────────────────────────────────────────────────

@role_required('member')
def member_dashboard(request):
    profile = request.user.userprofile
    org = profile.organization
    open_cycle = AwardCycle.objects.filter(is_open=True).first()
    questionnaire = None
    if open_cycle and org:
        questionnaire = Questionnaire.objects.filter(cycle=open_cycle, organization=org).first()
    return render(request, 'assessment/dashboard_member.html', {
        'org': org,
        'open_cycle': open_cycle,
        'questionnaire': questionnaire,
    })


@role_required('member')
def member_questionnaire_redirect(request):
    profile = request.user.userprofile
    org = profile.organization
    open_cycle = AwardCycle.objects.filter(is_open=True).first()
    if open_cycle and org:
        questionnaire = Questionnaire.objects.filter(cycle=open_cycle, organization=org).first()
        if questionnaire:
            if questionnaire.is_submitted:
                return redirect('questionnaire_submitted', pk=questionnaire.pk)
            return redirect('questionnaire_fill', pk=questionnaire.pk)
    messages.warning(request, 'No active questionnaire found.')
    return redirect('member_dashboard')


@role_required('member')
def questionnaire_fill(request, pk):
    questionnaire = get_object_or_404(Questionnaire, pk=pk)
    profile = request.user.userprofile
    if questionnaire.organization != profile.organization:
        raise PermissionDenied
    if questionnaire.is_submitted:
        return redirect('questionnaire_submitted', pk=pk)
    if not questionnaire.cycle.is_open:
        messages.warning(request, 'This cycle is closed.')
        return redirect('member_dashboard')

    category_id = request.GET.get('category')
    categories = _categories_for_questionnaire(questionnaire)
    current_category = (
        get_object_or_404(
            categories, id=category_id,
        ) if category_id
        else categories.first()
    )
    categories_list = list(
        categories.prefetch_related(
            Prefetch('criteria', queryset=Criterion.objects.filter(is_active=True), to_attr='active_criteria')
        )
    )
    current_category_index = next(
        (i + 1 for i, c in enumerate(categories_list) if c.id == current_category.id), 1
    ) if current_category else 1

    answered_ids = set(
        questionnaire.responses.exclude(response='').values_list('criterion_id', flat=True)
    )
    total_criteria = 0
    for cat in categories_list:
        cat_criteria_ids = {c.id for c in cat.active_criteria}
        total_criteria += len(cat_criteria_ids)
        cat.is_complete = bool(cat_criteria_ids) and cat_criteria_ids.issubset(answered_ids)
    answered_count = len(answered_ids)

    org_cycles = list(
        AwardCycle.objects
        .filter(questionnaires__organization=questionnaire.organization)
        .order_by('-year')
        .values('id', 'name')
    )
    # A questionnaire must have at least one applicable category.  This can
    # legitimately be empty for an incorrectly tagged/imported cycle, but it
    # should never crash with AttributeError.
    if current_category is None:
        messages.error(
            request,
            'No assessment sections are configured for this organisation in the current cycle. '
            'Please ask the administrator to import/check the questionnaire sections.'
        )
        return redirect('member_dashboard')

    form = BulkResponseForm(category=current_category, questionnaire=questionnaire)
    criteria_data = []
    for criterion in current_category.criteria.filter(is_active=True).order_by('order', 'number'):
        existing = Response.objects.filter(questionnaire=questionnaire, criterion=criterion).first()
        links = list(existing.evidence_links.select_related('document')) if existing else []
        criteria_data.append({
            'criterion': criterion,
            'levels': criterion.levels.order_by('level'),
            'existing_response': existing,
            'links': links,
            'numeric_values': [],
        })

    # Build the hierarchy used by the member entry page:
    # Assessment Category (current_category) -> Assessment Area ->
    # Assessment Criteria -> measurable parameters (Criterion records).
    from collections import OrderedDict
    area_groups_map = OrderedDict()
    for item in criteria_data:
        criterion = item['criterion']
        area_code = (criterion.area_code or '').strip()
        area_name = (criterion.assessment_area or '').strip() or 'Other / Unclassified Area'
        area_key = (area_code, area_name)
        if area_key not in area_groups_map:
            area_groups_map[area_key] = {
                'code': area_code,
                'name': area_name,
                'criteria_groups': OrderedDict(),
            }
        assessment_criterion_code = (criterion.assessment_criterion_code or '').strip()
        assessment_criterion_name = (criterion.assessment_criterion or '').strip() or 'Other / Unclassified Criterion'
        criterion_key = (assessment_criterion_code, assessment_criterion_name)
        criterion_groups = area_groups_map[area_key]['criteria_groups']
        if criterion_key not in criterion_groups:
            criterion_groups[criterion_key] = {
                'code': assessment_criterion_code,
                'name': assessment_criterion_name,
                'items': [],
            }
        criterion_groups[criterion_key]['items'].append(item)

    area_groups = [
        {
            'code': area['code'],
            'name': area['name'],
            'criteria_groups': list(area['criteria_groups'].values()),
        }
        for area in area_groups_map.values()
    ]

    from django.conf import settings
    from django.template.defaultfilters import filesizeformat as _filesizeformat
    quota_used = (
        EvidenceDocument.objects
        .filter(organization=questionnaire.organization)
        .aggregate(t=Sum('file_size'))['t'] or 0
    )
    quota_max = settings.EVIDENCE_MAX_QUOTA_BYTES
    quota_pct = min(round(quota_used / quota_max * 100), 100) if quota_max else 0

    return render(request, 'assessment/questionnaire_fill.html', {
        'questionnaire': questionnaire,
        'categories': categories_list,
        'current_category': current_category,
        'form': form,
        'criteria_data': criteria_data,
        'area_groups': area_groups,
        'current_category_index': current_category_index,
        'answered_count': answered_count,
        'total_criteria': total_criteria,
        'quota_used': quota_used,
        'quota_pct': quota_pct,
        'quota_max_display': _filesizeformat(quota_max),
        'org_cycles': org_cycles,
        'max_file_bytes': settings.EVIDENCE_MAX_FILE_BYTES,
        'max_file_display': _filesizeformat(settings.EVIDENCE_MAX_FILE_BYTES),
    })


@role_required('member')
def save_category(request, pk):
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Method not allowed.'}, status=405)

    questionnaire = get_object_or_404(Questionnaire, pk=pk)
    profile = request.user.userprofile
    if questionnaire.organization != profile.organization:
        raise PermissionDenied
    if questionnaire.is_submitted:
        return JsonResponse({'success': False, 'error': 'Questionnaire already submitted.'}, status=403)
    if not questionnaire.cycle.is_open:
        return JsonResponse({'success': False, 'error': 'Cycle is closed.'}, status=403)

    cat_id = request.POST.get('category_id')
    category = get_object_or_404(
        _categories_for_questionnaire(questionnaire), pk=cat_id
    )
    form = BulkResponseForm(request.POST, category=category, questionnaire=questionnaire)
    if not form.is_valid():
        return JsonResponse({'success': False, 'errors': form.errors}, status=400)

    with transaction.atomic():
        for criterion in category.criteria.filter(is_active=True):
            response_text = (form.cleaned_data.get(f'response_{criterion.id}', '') or '').strip()
            if response_text:
                Response.objects.update_or_create(
                    questionnaire=questionnaire, criterion=criterion,
                    defaults={'response': response_text, 'score': None, 'numeric_data': None},
                )
            else:
                Response.objects.filter(questionnaire=questionnaire, criterion=criterion).update(
                    response='', score=None, numeric_data=None
                )

    categories = _categories_for_questionnaire(questionnaire)
    next_cat = categories.filter(order__gt=category.order).first()
    next_url = None
    if next_cat:
        next_url = reverse('questionnaire_fill', args=[pk]) + f'?category={next_cat.id}'

    return JsonResponse({
        'success': True,
        'message': f'Saved: {category.name}',
        'next_url': next_url,
    })


@role_required('member')
def evidence_upload_ajax(request):
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Method not allowed.'}, status=405)

    profile = request.user.userprofile
    q = get_object_or_404(Questionnaire, pk=request.POST.get('questionnaire_id'))
    if q.organization_id != profile.organization_id:
        raise PermissionDenied
    criterion = get_object_or_404(
        Criterion.objects.filter(
            category__in=_categories_for_questionnaire(q),
            category__is_active=True,
            is_active=True,
        ),
        pk=request.POST.get('criterion_id'),
    )
    if q.is_submitted or not q.cycle.is_open:
        return JsonResponse({'success': False, 'error': 'Locked.'}, status=403)

    f = request.FILES.get('file')
    if not f:
        return JsonResponse({'success': False, 'error': 'No file provided.'}, status=400)

    resp, _ = Response.objects.get_or_create(questionnaire=q, criterion=criterion)
    err = _validate_evidence_upload(f, q, resp)
    if err:
        return JsonResponse({'success': False, 'error': err}, status=400)

    safe_name = _re.sub(r'[^\w.\-]', '_', _os.path.basename(f.name))
    title = _os.path.splitext(safe_name)[0]
    evidence_doc = EvidenceDocument(
        organization=q.organization,
        title=title,
        original_filename=f.name,
        file_size=f.size,
        uploaded_by=request.user,
    )
    path = f'evidence/{q.organization_id}/{_uuid.uuid4().hex}_{safe_name}'
    evidence_doc.file.save(path, f, save=True)
    link = EvidenceLink.objects.create(response=resp, document=evidence_doc)

    return JsonResponse({
        'success': True,
        'document': {
            'id': evidence_doc.id,
            'title': evidence_doc.title,
            'original_filename': evidence_doc.original_filename,
            'file_url': reverse('document_file', args=[evidence_doc.id]),
            'delete_url': reverse('evidence_unlink', args=[link.id]),
        },
    })


@role_required('member')
def questionnaire_submit(request, pk):
    questionnaire = get_object_or_404(Questionnaire, pk=pk)
    profile = request.user.userprofile
    if questionnaire.organization != profile.organization:
        raise PermissionDenied
    if questionnaire.is_submitted:
        return redirect('questionnaire_submitted', pk=pk)
    if request.method == 'POST':
        answered = questionnaire.responses.exclude(response='').count()
        if answered == 0:
            messages.error(request, 'Please enter at least one Response before submitting.')
            return redirect('questionnaire_fill', pk=pk)
        questionnaire.is_submitted = True
        questionnaire.submitted_at = timezone.now()
        questionnaire.save()
        AuditLog.log(request, 'questionnaire.submitted',
            f'Submitted questionnaire for "{questionnaire.cycle.name}" — {questionnaire.organization.name}',
            object_type='Questionnaire', object_id=questionnaire.pk,
            object_repr=str(questionnaire), organization=questionnaire.organization)
        messages.success(request, 'Questionnaire submitted successfully!')
        return redirect('questionnaire_submitted', pk=pk)
    return redirect('questionnaire_fill', pk=pk)


@role_required('member')
def questionnaire_submitted(request, pk):
    questionnaire = get_object_or_404(Questionnaire, pk=pk)
    profile = request.user.userprofile
    if questionnaire.organization != profile.organization:
        raise PermissionDenied
    if not questionnaire.is_submitted:
        return redirect('questionnaire_fill', pk=pk)
    categories = list(
        _categories_for_questionnaire(questionnaire).prefetch_related(
            Prefetch('criteria', queryset=Criterion.objects.filter(is_active=True).order_by('number'),
                     to_attr='active_criteria')
        )
    )
    resp_map = {r.criterion_id: r for r in questionnaire.responses.all()}
    cat_data = []
    for cat in categories:
        rows = [
            {'criterion': c, 'response': resp_map.get(c.id)}
            for c in cat.active_criteria
        ]
        cat_data.append({'category': cat, 'rows': rows})
    return render(request, 'assessment/questionnaire_submitted.html', {
        'questionnaire': questionnaire, 'cat_data': cat_data,
    })



@role_required('member')
def evidence_link(request):
    from django.http import JsonResponse
    from django.conf import settings
    if request.method != 'POST':
        return JsonResponse({'ok': False}, status=405)
    profile = request.user.userprofile
    doc = get_object_or_404(EvidenceDocument, pk=request.POST.get('document_id') or request.POST.get('doc_id'))

    # Two entry points: an existing Response (legacy) or a questionnaire+criterion
    # pair. The latter lets a member attach evidence before scoring the criterion,
    # so the library picker works on a fresh cycle's empty questionnaire.
    response_id = request.POST.get('response_id')
    if response_id:
        resp = get_object_or_404(Response, pk=response_id)
        q = resp.questionnaire
        criterion = None
    else:
        q = get_object_or_404(Questionnaire, pk=request.POST.get('questionnaire_id'))
        criterion = get_object_or_404(
            Criterion.objects.filter(
                category__in=_categories_for_questionnaire(q),
                category__is_active=True,
                is_active=True,
            ),
            pk=request.POST.get('criterion_id'),
        )
        resp = None

    if q.organization_id != profile.organization_id:
        return JsonResponse({'ok': False, 'error': 'Forbidden.'}, status=403)
    if doc.organization_id != profile.organization_id:
        return JsonResponse({'ok': False, 'error': 'Forbidden.'}, status=403)
    if q.is_submitted or not q.cycle.is_open:
        return JsonResponse({'ok': False, 'error': 'Locked.'}, status=403)

    if resp is None:
        resp, _ = Response.objects.get_or_create(questionnaire=q, criterion=criterion)
    if resp.evidence_links.count() >= settings.EVIDENCE_MAX_FILES_PER_RESP:
        return JsonResponse({'ok': False, 'error': 'Max links per criterion.'}, status=409)
    EvidenceLink.objects.get_or_create(response=resp, document=doc)
    return JsonResponse({'ok': True})


@role_required('member')
def evidence_unlink(request, link_pk):
    from django.http import JsonResponse
    if request.method != 'POST':
        return JsonResponse({'ok': False}, status=405)
    link = get_object_or_404(EvidenceLink, pk=link_pk)
    q = link.response.questionnaire
    profile = request.user.userprofile
    if q.organization_id != profile.organization_id:
        return JsonResponse({'ok': False, 'error': 'Forbidden.'}, status=403)
    if q.is_submitted or not q.cycle.is_open:
        return JsonResponse({'ok': False, 'error': 'Locked.'}, status=403)
    document = link.document
    link.delete()
    if not q.is_submitted and document.evidence_links.count() == 0:
        document.delete()  # django-cleanup removes the file from storage
    return JsonResponse({'ok': True})


# ── Library ───────────────────────────────────────────────────────────────────

def _resolve_library_org(request):
    """Member → own org; admin → org from ?org=<id>."""
    profile = request.user.userprofile
    if profile.role == 'member':
        return profile.organization
    if profile.role == 'admin':
        org_id = request.GET.get('org') or request.POST.get('org')
        return get_object_or_404(Organization, pk=org_id) if org_id else None
    raise PermissionDenied


@login_required
def library_list(request):
    profile = request.user.userprofile
    if profile.role not in ('member', 'admin'):
        raise PermissionDenied
    org = _resolve_library_org(request)
    documents = []
    quota_used = 0
    if org:
        documents = list(
            EvidenceDocument.objects
            .filter(organization=org)
            .prefetch_related(
                Prefetch(
                    'evidence_links',
                    queryset=EvidenceLink.objects
                        .select_related(
                            'response__criterion',
                            'response__questionnaire__cycle',
                        )
                        .order_by('created_at'),
                    to_attr='all_links',
                )
            )
            .order_by('-uploaded_at')
        )
        quota_used = (
            EvidenceDocument.objects
            .filter(organization=org)
            .aggregate(t=Sum('file_size'))['t'] or 0
        )
    from django.conf import settings
    from django.template.defaultfilters import filesizeformat
    return render(request, 'assessment/library_list.html', {
        'org': org,
        'documents': documents,
        'quota_used': quota_used,
        'quota_max': settings.EVIDENCE_MAX_QUOTA_BYTES,
        'all_orgs': Organization.objects.filter(is_active=True) if profile.role == 'admin' else None,
    })


# ── Verifier ─────────────────────────────────────────────────────────────────

@role_required('verifier')
def verifier_dashboard(request):
    from accounts.models import VerifierAssignment
    assignments = VerifierAssignment.objects.filter(
        verifier=request.user
    ).select_related('organization')
    open_cycle = AwardCycle.objects.filter(is_open=True).first()

    questionnaires = []
    q_by_org = {}
    if open_cycle:
        questionnaires = list(
            Questionnaire.objects
            .filter(cycle=open_cycle, organization__in=[a.organization for a in assignments])
            .select_related('organization')
        )
        q_by_org = {q.organization_id: q for q in questionnaires}

    coverage_map = questionnaire_coverage(
        [q for q in questionnaires if q.is_submitted]
    )

    org_data = []
    for a in assignments:
        q = q_by_org.get(a.organization_id)
        org_data.append({
            'org': a.organization,
            'questionnaire': q,
            'coverage': coverage_map.get(q.pk) if q and q.is_submitted else None,
        })

    submitted_count = 0
    pending_count = 0
    if open_cycle:
        submitted_count = sum(
            1 for item in org_data
            if item['questionnaire'] and item['questionnaire'].is_submitted
        )
        pending_count = len(org_data) - submitted_count

    return render(request, 'assessment/dashboard_verifier.html', {
        'org_data': org_data,
        'open_cycle': open_cycle,
        'assigned_orgs': assignments,
        'submitted_count': submitted_count,
        'pending_count': pending_count,
    })




def _refresh_judging_completion(questionnaire):
    """Mark judging complete only after every active judge has submitted.

    Each judge must have a submitted StageSubmission and a score for every
    applicable non-numeric criterion. Missing judges/scores are never treated
    as zero.
    """
    available_judge_ids = set(User.objects.filter(
        userprofile__role='judge', is_active=True
    ).values_list('id', flat=True))
    if not available_judge_ids:
        questionnaire.judging_completed = False
        questionnaire.save(update_fields=['judging_completed', 'updated_at'])
        return False

    applicable_category_ids = set(_categories_for_questionnaire(questionnaire).values_list('id', flat=True))
    criterion_ids = set(Criterion.objects.filter(
        category__id__in=applicable_category_ids,
        category__is_active=True, is_active=True, is_numeric=False,
    ).values_list('id', flat=True))
    if not criterion_ids:
        questionnaire.judging_completed = False
        questionnaire.save(update_fields=['judging_completed', 'updated_at'])
        return False

    submitted_judge_ids = set(StageSubmission.objects.filter(
        questionnaire=questionnaire, stage='judge', category__isnull=True,
        user_id__in=available_judge_ids, submitted_at__isnull=False,
    ).values_list('user_id', flat=True))

    if submitted_judge_ids != available_judge_ids:
        questionnaire.judging_completed = False
        questionnaire.save(update_fields=['judging_completed', 'updated_at'])
        return False

    actual_pairs = set(JudgeResponse.objects.filter(
        questionnaire=questionnaire,
        judge_id__in=available_judge_ids,
        criterion_id__in=criterion_ids,
        score__isnull=False,
    ).values_list('judge_id', 'criterion_id'))
    expected_pairs = {(judge_id, criterion_id)
                      for judge_id in available_judge_ids
                      for criterion_id in criterion_ids}
    complete = actual_pairs >= expected_pairs
    questionnaire.judging_completed = complete
    if complete:
        questionnaire.judged_at = timezone.now()
        questionnaire.judged_by = None
        questionnaire.save(update_fields=['judging_completed', 'judged_at', 'judged_by', 'updated_at'])
    else:
        questionnaire.save(update_fields=['judging_completed', 'updated_at'])
    return complete


@role_required('judge')
def judges_dashboard(request):
    """Dashboard for judges. Each judge works independently."""
    open_cycle = AwardCycle.objects.filter(is_open=True).first()
    org_data = []
    pending_count = 0
    judged_count = 0
    available_judges = User.objects.filter(
        userprofile__role='judge', is_active=True
    ).count()

    if open_cycle:
        questionnaires = list(
            Questionnaire.objects
            .filter(
                cycle=open_cycle,
                organization__requires_judging=True,
                is_submitted=True,
            )
            .select_related('organization')
        )
        coverage_map = questionnaire_coverage(questionnaires)

        for q in questionnaires:
            covered, total = coverage_map.get(q.pk, (0, 0))
            secretariat_complete = total > 0 and covered >= total
            scorable_ids = list(
                Criterion.objects.filter(
                    category__cycle=q.cycle,
                    category__is_active=True,
                    is_active=True,
                    is_numeric=False,
                ).filter(
                    category_id__in=_categories_for_questionnaire(q).values_list('id', flat=True)
                ).values_list('id', flat=True)
            )
            my_scored = JudgeResponse.objects.filter(
                questionnaire=q, judge=request.user,
                criterion_id__in=scorable_ids, score__isnull=False,
            ).count()
            my_complete = StageSubmission.objects.filter(
                questionnaire=q, user=request.user, stage='judge', category__isnull=True,
                submitted_at__isnull=False
            ).exists()
            submitted_judges = StageSubmission.objects.filter(
                questionnaire=q, stage='judge', category__isnull=True, submitted_at__isnull=False
            ).values('user_id').distinct().count()

            if not secretariat_complete:
                status = 'waiting'
            elif my_complete:
                status = 'judged'
                judged_count += 1
            else:
                status = 'ready'
                pending_count += 1

            org_data.append({
                'questionnaire': q,
                'org': q.organization,
                'covered': covered,
                'total': total,
                'verifier_complete': secretariat_complete,
                'status': status,
                'my_complete': my_complete,
                'submitted_judges': submitted_judges,
                'available_judges': available_judges,
            })

    return render(request, 'assessment/dashboard_judges.html', {
        'open_cycle': open_cycle,
        'org_data': org_data,
        'pending_count': pending_count,
        'judged_count': judged_count,
        'available_judges': available_judges,
    })


@role_required('judge')
def judge_questionnaire(request, pk):
    """Judge one questionnaire without exposing any other judge's scoring."""
    questionnaire = get_object_or_404(
        Questionnaire.objects.select_related('organization', 'cycle'),
        pk=pk,
        is_submitted=True,
        organization__requires_judging=True,
    )

    if not questionnaire.cycle.is_open:
        messages.info(request, 'This cycle is closed. Judge scores are locked.')
        return redirect('judges_dashboard')

    coverage = questionnaire_coverage([questionnaire]).get(questionnaire.pk, (0, 0))
    covered, total = coverage
    secretariat_complete = total > 0 and covered >= total

    if not secretariat_complete:
        messages.warning(
            request,
            'This organization is not ready for judging yet. The Secretariat must complete all applicable criteria first.'
        )
        return redirect('judges_dashboard')

    applicable_categories = _categories_for_questionnaire(questionnaire)
    applicable_category_ids = set(applicable_categories.values_list('id', flat=True))
    criteria = list(
        Criterion.objects.filter(
            category__cycle=questionnaire.cycle,
            category__is_active=True,
            is_active=True,
            category_id__in=applicable_category_ids,
        ).prefetch_related('levels').order_by(
            'category__order', 'category__name', 'order', 'number'
        )
    )
    crit_ids = [c.id for c in criteria]
    scorable_ids = [c.id for c in criteria if not c.is_numeric]

    member_map = {
        r.criterion_id: r for r in
        Response.objects.filter(questionnaire=questionnaire, criterion_id__in=crit_ids)
        .prefetch_related('evidence_links__document')
    }
    secretariat_map = {}
    for sr in (VerifierResponse.objects.filter(
        questionnaire=questionnaire, criterion_id__in=crit_ids
    ).select_related('verifier').order_by('criterion_id', 'created_at')):
        secretariat_map.setdefault(sr.criterion_id, []).append(sr)

    # SECURITY: only this judge's own rows are loaded. No queryset containing
    # another judge's score/comment is ever passed to this template.
    my_judge_map = {
        jr.criterion_id: jr for jr in JudgeResponse.objects.filter(
            questionnaire=questionnaire,
            judge=request.user,
            criterion_id__in=crit_ids,
        )
    }

    # A submitted judge assessment is locked while the cycle remains open.
    judge_submission, _ = StageSubmission.objects.get_or_create(
        questionnaire=questionnaire, user=request.user, stage='judge', category=None
    )

    if request.method == 'POST':
        action = request.POST.get('action', 'save')
        if judge_submission.is_submitted:
            messages.warning(request, 'Your Judge assessment has already been submitted and is locked.')
            return redirect('judge_questionnaire', pk=pk)

        saved_count = 0
        with transaction.atomic():
            for criterion in criteria:
                score_raw = request.POST.get(f'judge_score_{criterion.id}', '').strip()
                notes = request.POST.get(f'judge_notes_{criterion.id}', '').strip()
                excel_file = request.FILES.get(f'judge_excel_{criterion.id}')
                score = None
                if not criterion.is_numeric and score_raw:
                    try:
                        score = int(score_raw)
                    except (ValueError, TypeError):
                        score = None
                    if score is not None and not 0 <= score <= 5:
                        score = None

                excel_error = _validate_judge_excel_upload(excel_file)
                if excel_error:
                    messages.error(request, excel_error)
                    return redirect('judge_questionnaire', pk=pk)

                existing = JudgeResponse.objects.filter(
                    questionnaire=questionnaire,
                    judge=request.user,
                    criterion=criterion,
                ).first()
                if score is None and not notes and not excel_file and not existing:
                    continue

                obj, _ = JudgeResponse.objects.update_or_create(
                    questionnaire=questionnaire,
                    judge=request.user,
                    criterion=criterion,
                    defaults={'score': score, 'notes': notes},
                )
                if excel_file:
                    safe_name = _re.sub(r'[^\\w.\\-]', '_', _os.path.basename(excel_file.name))
                    obj.excel_attachment.save(
                        f'{_uuid.uuid4().hex}_{safe_name}', excel_file, save=True
                    )
                saved_count += 1

            if action == 'submit':
                # Submit means the judge is declaring this questionnaire complete.
                # Require every non-numeric criterion to have a score.
                missing = JudgeResponse.objects.filter(
                    questionnaire=questionnaire, judge=request.user,
                    criterion_id__in=scorable_ids, score__isnull=True,
                ).count()
                existing_scored = JudgeResponse.objects.filter(
                    questionnaire=questionnaire, judge=request.user,
                    criterion_id__in=scorable_ids, score__isnull=False,
                ).count()
                if existing_scored < len(scorable_ids) or missing:
                    messages.error(
                        request,
                        'Please enter a score for every applicable criterion before submitting.'
                    )
                    return redirect('judge_questionnaire', pk=pk)
                judge_submission.submitted_at = timezone.now()
                judge_submission.save(update_fields=['submitted_at', 'updated_at'])

        if saved_count:
            AuditLog.log(
                request, 'judge.scores_saved',
                f'Saved judge scores for "{questionnaire.organization.name}" — {questionnaire.cycle.name}',
                object_type='Questionnaire', object_id=questionnaire.pk,
                object_repr=str(questionnaire), organization=questionnaire.organization,
            )

        if action == 'submit':
            _refresh_judging_completion(questionnaire)
            messages.success(request, 'Your Judge assessment has been submitted successfully. It is now locked.')
            return redirect('judges_dashboard')
        if action == 'save_edit':
            messages.success(request, f'Saved {saved_count} judging response(s). You can continue editing.')
            return redirect('judge_questionnaire', pk=pk)
        if saved_count:
            messages.success(request, f'Saved {saved_count} judging response(s).')
        else:
            messages.warning(request, 'No judging responses were entered.')
        return redirect('judges_dashboard')

    # Re-read after POST or on GET.
    my_judge_map = {
        jr.criterion_id: jr for jr in JudgeResponse.objects.filter(
            questionnaire=questionnaire, judge=request.user, criterion_id__in=crit_ids
        )
    }

    criteria_data = []
    for criterion in criteria:
        member_resp = member_map.get(criterion.id)
        criteria_data.append({
            'criterion': criterion,
            'levels': criterion.levels.all(),
            'member_response': member_resp,
            'secretariat_responses': secretariat_map.get(criterion.id, []),
            'judge_response': my_judge_map.get(criterion.id),
            'links': member_resp.evidence_links.all() if member_resp else [],
        })

    scorable_ids = [c.id for c in criteria if not c.is_numeric]
    my_scored = JudgeResponse.objects.filter(
        questionnaire=questionnaire, judge=request.user,
        criterion_id__in=scorable_ids, score__isnull=False,
    ).count()
    my_complete = judge_submission.is_submitted
    available_judges = User.objects.filter(
        userprofile__role='judge', is_active=True
    ).count()
    submitted_judges = StageSubmission.objects.filter(
        questionnaire=questionnaire, stage='judge', category__isnull=True,
        submitted_at__isnull=False
    ).values('user_id').distinct().count()

    return render(request, 'assessment/judge_questionnaire.html', {
        'questionnaire': questionnaire,
        'criteria_data': criteria_data,
        'verifier_complete': secretariat_complete,
        'covered': covered,
        'total': total,
        'my_complete': my_complete,
        'available_judges': available_judges,
        'submitted_judges': submitted_judges,
        'judge_submitted': judge_submission.is_submitted,
    })

@role_required('verifier')
def verify_questionnaire(request, pk):
    questionnaire = get_object_or_404(Questionnaire, pk=pk, is_submitted=True)
    from accounts.models import VerifierAssignment
    assigned = VerifierAssignment.objects.filter(
        verifier=request.user, organization=questionnaire.organization
    ).exists()
    if not assigned:
        raise PermissionDenied
    if not questionnaire.cycle.is_open:
        messages.warning(request, 'This cycle is closed — scores are locked.')

    is_informal = questionnaire.organization.org_type == 'informal_sector'
    category_id = request.GET.get('category')
    categories = _categories_for_questionnaire(questionnaire)
    current_category = (
        get_object_or_404(
            categories, id=category_id,
        ) if category_id
        else categories.first()
    )

    secretariat_submission = None
    if current_category:
        secretariat_submission, _ = StageSubmission.objects.get_or_create(
            questionnaire=questionnaire, user=request.user,
            stage='secretariat', category=current_category
        )

    if request.method == 'POST' and questionnaire.cycle.is_open:
        action = request.POST.get('action', 'save')
        cat_id = request.POST.get('category_id')
        category = get_object_or_404(
            _categories_for_questionnaire(questionnaire), pk=cat_id
        )
        category_submission, _ = StageSubmission.objects.get_or_create(
            questionnaire=questionnaire, user=request.user,
            stage='secretariat', category=category
        )
        if category_submission.is_submitted:
            messages.warning(request, 'This Secretariat section has already been submitted and is locked.')
            return redirect(f'{request.path}?category={category.id}')

        form = VerifierScoreForm(
            request.POST, category=category,
            questionnaire=questionnaire, verifier=request.user
        )
        if not form.is_valid():
            messages.error(request, 'Please correct the errors below.')
            return redirect(f'{request.path}?category={category.id}')

        saved_count = 0
        with transaction.atomic():
            for criterion in category.criteria.filter(is_active=True):
                notes = form.cleaned_data.get(f'notes_{criterion.id}', '') or ''

                # Secretariat is never allowed to create/update a score.
                if notes.strip():
                    VerifierResponse.objects.update_or_create(
                        questionnaire=questionnaire,
                        verifier=request.user,
                        criterion=criterion,
                        defaults={'score': None, 'notes': notes},
                    )
                    saved_count += 1

                # Secretariat may attach evidence to this exact criterion even
                # though the member questionnaire has already been submitted.
                files_for_criterion = request.FILES.getlist(
                    f'files_{criterion.id}'
                )
                response_obj, _ = Response.objects.get_or_create(
                    questionnaire=questionnaire, criterion=criterion
                )
                for uploaded in files_for_criterion:
                    err = _validate_evidence_upload(
                        uploaded, questionnaire, response_obj
                    )
                    if err:
                        messages.error(request, err)
                        transaction.set_rollback(True)
                        return redirect(f'{request.path}?category={category.id}')

                    safe_name = _re.sub(
                        r'[^\\w.\\-]', '_',
                        _os.path.basename(uploaded.name)
                    )
                    title = _os.path.splitext(safe_name)[0]
                    evidence_doc = EvidenceDocument(
                        organization=questionnaire.organization,
                        title=title,
                        original_filename=uploaded.name,
                        file_size=uploaded.size,
                        uploaded_by=request.user,
                    )
                    path = (
                        f'evidence/{questionnaire.organization_id}/'
                        f'{_uuid.uuid4().hex}_{safe_name}'
                    )
                    evidence_doc.file.save(path, uploaded, save=True)
                    EvidenceLink.objects.create(
                        response=response_obj, document=evidence_doc
                    )
                    saved_count += 1

            if action == 'submit':
                # Secretariat submits a category to declare its review/evidence
                # work complete. No score is required or accepted.
                category_submission.submitted_at = timezone.now()
                category_submission.save(
                    update_fields=['submitted_at', 'updated_at']
                )

        if saved_count:
            AuditLog.log(
                request, 'verifier.review_saved',
                f'Saved Secretariat review for "{category.name}" — {questionnaire.organization.name}',
                object_type='Questionnaire', object_id=questionnaire.pk,
                object_repr=str(questionnaire), organization=questionnaire.organization
            )

        if action == 'submit':
            messages.success(request, f'Submitted: {category.name}. This section is now locked.')
            next_cat = categories.filter(order__gt=category.order).first()
            if next_cat:
                return redirect(f'{request.path}?category={next_cat.id}')
            return redirect('verifier_dashboard')
        if action == 'save_edit':
            messages.success(request, f'Saved: {category.name}. You can continue editing.')
            return redirect(f'{request.path}?category={category.id}')

        messages.success(request, f'Saved: {category.name}.')
        return redirect('verifier_dashboard')

    form = VerifierScoreForm(
        category=current_category, questionnaire=questionnaire, verifier=request.user
    ) if current_category else None
    criteria = list(
        current_category.criteria.filter(is_active=True).order_by('number').prefetch_related('levels')
        if current_category else []
    )
    crit_ids = [c.id for c in criteria]
    member_map = {
        r.criterion_id: r for r in
        Response.objects.filter(questionnaire=questionnaire, criterion_id__in=crit_ids)
        .prefetch_related('evidence_links__document')
    }
    verifier_map = {
        vr.criterion_id: vr for vr in
        VerifierResponse.objects.filter(
            questionnaire=questionnaire, verifier=request.user, criterion_id__in=crit_ids
        )
    }
    judge_map = {}
    if not questionnaire.cycle.is_open:
        for jr in JudgeResponse.objects.filter(
            questionnaire=questionnaire, criterion_id__in=crit_ids
        ).select_related('judge').order_by('criterion_id', 'judge__first_name', 'judge__username'):
            if jr.score is not None or (jr.notes or '').strip():
                judge_map.setdefault(jr.criterion_id, []).append(jr)

    coverage = criterion_coverage(questionnaire, criteria)
    cat_coverage = category_coverage(questionnaire)

    criteria_data = []
    for criterion in criteria:
        member_resp = member_map.get(criterion.id)
        others = [
            entry for entry in coverage.get(criterion.id, [])
            if entry['verifier_id'] != request.user.id
        ]
        scores = [r.score for r in judge_map.get(criterion.id, []) if r.score is not None]
        criteria_data.append({
            'criterion': criterion,
            'levels': criterion.levels.all(),
            'member_response': member_resp,
            'verifier_response': verifier_map.get(criterion.id),
            'judge_responses': judge_map.get(criterion.id, []),
            'judge_average': round(sum(scores) / len(scores), 1) if scores else None,
            'links': member_resp.evidence_links.all() if member_resp else [],
            'others': others,
            'others_extra': len(others) - 1 if others else 0,
            'others_title': '; '.join(
                f"{entry['name']} ({entry['at'].strftime('%d %b %Y')})"
                for entry in others
            ),
        })

    category_nav = []
    for cat in categories:
        covered, total = cat_coverage.get(cat.id, (0, 0))
        submitted = StageSubmission.objects.filter(
            questionnaire=questionnaire, user=request.user, stage='secretariat',
            category=cat, submitted_at__isnull=False
        ).exists()
        category_nav.append({
            'category': cat, 'covered': covered, 'total': total, 'submitted': submitted
        })

    return render(request, 'assessment/verify_questionnaire.html', {
        'questionnaire': questionnaire,
        'categories': categories,
        'category_nav': category_nav,
        'current_category': current_category,
        'form': form,
        'criteria_data': criteria_data,
        'cycle_open': questionnaire.cycle.is_open,
        'secretariat_submitted': secretariat_submission.is_submitted if secretariat_submission else False,
    })


@login_required
def document_file(request, doc_pk):
    from assessment.access import can_access_document
    doc = get_object_or_404(EvidenceDocument, pk=doc_pk)
    if not can_access_document(request.user, doc):
        raise PermissionDenied
    download = bool(request.GET.get('download'))
    storage = doc.file.storage
    name = doc.file.name
    disposition = f'attachment; filename="{doc.original_filename}"'

    if storage.__class__.__name__ == 'S3Boto3Storage':
        if download:
            return HttpResponseRedirect(storage.url(
                name, parameters={'ResponseContentDisposition': disposition}))
        return HttpResponseRedirect(doc.file.url)

    if not name:
        return redirect('library_list')

    if download and storage.exists(name):
        from django.http import FileResponse
        return FileResponse(
            doc.file.open('rb'), as_attachment=True,
            filename=doc.original_filename)
    return HttpResponseRedirect(doc.file.url)


@login_required
def library_search(request):
    from django.http import JsonResponse
    from django.db.models import Q
    profile = request.user.userprofile
    if profile.role not in ('member', 'admin'):
        return JsonResponse({'error': 'Forbidden'}, status=403)

    if profile.role == 'member':
        org = profile.organization
    else:
        org_id = request.GET.get('org')
        if not org_id:
            return JsonResponse({'results': []})
        org = get_object_or_404(Organization, pk=org_id)

    q_text = request.GET.get('q', '').strip()
    cycle_filter = request.GET.get('cycle', '').strip()
    criterion_id = request.GET.get('criterion_id', '').strip()

    qs = EvidenceDocument.objects.filter(organization=org)

    if q_text:
        qs = qs.filter(
            Q(title__icontains=q_text) | Q(original_filename__icontains=q_text)
        )

    if cycle_filter:
        qs = qs.filter(
            evidence_links__response__questionnaire__cycle_id=cycle_filter
        ).distinct()

    # Build ranking: same criterion → same category → everything else
    same_criterion_ids = set()
    same_category_ids = set()
    if criterion_id:
        same_criterion_ids = set(
            EvidenceLink.objects
            .filter(response__criterion_id=criterion_id, document__organization=org)
            .values_list('document_id', flat=True)
        )
        try:
            crit_obj = Criterion.objects.select_related('category').get(pk=criterion_id)
            same_category_ids = set(
                EvidenceLink.objects
                .filter(
                    response__criterion__category_id=crit_obj.category_id,
                    document__organization=org,
                )
                .values_list('document_id', flat=True)
            ) - same_criterion_ids
        except Criterion.DoesNotExist:
            pass

    all_docs = list(qs.distinct().order_by('-uploaded_at'))

    def rank(doc):
        if doc.id in same_criterion_ids:
            return 0
        if doc.id in same_category_ids:
            return 1
        return 2

    all_docs.sort(key=rank)
    top_docs = all_docs[:30]

    results = []
    for doc in top_docs:
        first_link = (
            EvidenceLink.objects
            .filter(document=doc)
            .select_related('response__criterion', 'response__questionnaire__cycle')
            .order_by('created_at')
            .first()
        )
        results.append({
            'id': doc.id,
            'title': doc.title,
            'original_filename': doc.original_filename,
            'uploaded_at': doc.uploaded_at.strftime('%d %b %Y'),
            'file_size': doc.file_size,
            'cycle_name': (
                first_link.response.questionnaire.cycle.name if first_link else ''
            ),
            'criterion_name': (
                f"{first_link.response.criterion.number} — {first_link.response.criterion.name}"
                if first_link else ''
            ),
            'is_suggested': doc.id in same_criterion_ids,
        })

    return JsonResponse({'results': results})


# ── Reports ──────────────────────────────────────────────────────────────────

@login_required
def reports_list(request):
    from accounts.models import VerifierAssignment
    profile = request.user.userprofile
    role = profile.role

    if role == 'admin':
        cycles = AwardCycle.objects.all()
        cycle_data = []
        for cycle in cycles:
            qs = cycle.questionnaires.all()
            cycle_data.append({
                'cycle': cycle,
                'org_count': qs.count(),
                'submitted_count': qs.filter(is_submitted=True).count(),
                'distributed_count': qs.filter(is_distributed=True).count(),
            })

    elif role == 'verifier':
        assigned_org_ids = VerifierAssignment.objects.filter(
            verifier=request.user
        ).values_list('organization_id', flat=True)
        cycles = AwardCycle.objects.filter(
            questionnaires__organization_id__in=assigned_org_ids
        ).distinct()
        cycle_data = []
        for cycle in cycles:
            qs = cycle.questionnaires.filter(organization_id__in=assigned_org_ids)
            cycle_data.append({
                'cycle': cycle,
                'org_count': qs.count(),
                'submitted_count': qs.filter(is_submitted=True).count(),
                'distributed_count': qs.filter(is_distributed=True).count(),
            })

    else:  # member
        org = profile.organization
        if not org:
            cycle_data = []
        else:
            questionnaires = (
                Questionnaire.objects
                .filter(organization=org)
                .select_related('cycle')
                .order_by('-cycle__year')
            )
            cycle_data = [
                {
                    'cycle': q.cycle,
                    'questionnaire': q,
                    'available': not q.cycle.is_open and q.is_distributed,
                }
                for q in questionnaires
            ]

    return render(request, 'assessment/reports_list.html', {
        'cycle_data': cycle_data,
        'role': role,
    })


@login_required
def reports_cycle(request, cycle_pk):
    from accounts.models import VerifierAssignment
    profile = request.user.userprofile
    role = profile.role
    cycle = get_object_or_404(AwardCycle, pk=cycle_pk)

    if role == 'admin':
        questionnaires = (
            cycle.questionnaires
            .select_related('organization', 'distributed_by')
            .order_by('organization__name')
        )
        return render(request, 'assessment/reports_cycle.html', {
            'cycle': cycle,
            'questionnaires': questionnaires,
            'role': role,
        })

    elif role == 'verifier':
        assigned_org_ids = VerifierAssignment.objects.filter(
            verifier=request.user
        ).values_list('organization_id', flat=True)
        questionnaires = (
            cycle.questionnaires
            .filter(organization_id__in=assigned_org_ids)
            .select_related('organization', 'distributed_by')
            .order_by('organization__name')
        )
        return render(request, 'assessment/reports_cycle.html', {
            'cycle': cycle,
            'questionnaires': questionnaires,
            'role': role,
        })

    else:  # member
        if not profile.organization:
            raise PermissionDenied
        questionnaire = get_object_or_404(
            Questionnaire, cycle=cycle, organization=profile.organization
        )
        if not cycle.is_open and questionnaire.is_distributed:
            return redirect('org_report', pk=cycle_pk, q_pk=questionnaire.pk)
        return render(request, 'assessment/reports_cycle.html', {
            'cycle': cycle,
            'questionnaire': questionnaire,
            'role': role,
        })


@login_required
def report_distribute(request, cycle_pk, q_pk):
    if request.method != 'POST':
        return redirect('reports_cycle', cycle_pk=cycle_pk)
    from accounts.models import VerifierAssignment
    profile = request.user.userprofile
    cycle = get_object_or_404(AwardCycle, pk=cycle_pk)
    questionnaire = get_object_or_404(Questionnaire, pk=q_pk, cycle=cycle)

    is_admin = profile.role == 'admin'
    is_assigned = (
        profile.role == 'verifier' and
        VerifierAssignment.objects.filter(
            verifier=request.user, organization=questionnaire.organization
        ).exists()
    )
    if not (is_admin or is_assigned):
        raise PermissionDenied

    if cycle.is_open:
        messages.error(request, 'Cannot distribute reports while the cycle is still open.')
        return redirect('reports_cycle', cycle_pk=cycle_pk)

    if questionnaire.organization.requires_judging and not questionnaire.judging_completed:
        messages.error(
            request,
            f'{questionnaire.organization.name} requires judging before its final results can be released.'
        )
        return redirect('reports_cycle', cycle_pk=cycle_pk)

    questionnaire.is_distributed = True
    questionnaire.distributed_at = timezone.now()
    questionnaire.distributed_by = request.user
    questionnaire.save(update_fields=['is_distributed', 'distributed_at', 'distributed_by'])

    AuditLog.log(
        request, 'report.distributed',
        f'Report distributed for {questionnaire.organization.name} ({cycle.name})',
        object_repr=str(questionnaire.organization),
        object_type='Questionnaire',
        object_id=questionnaire.pk,
        organization=questionnaire.organization,
    )
    messages.success(request, f'Report distributed to {questionnaire.organization.name}.')
    return redirect('reports_cycle', cycle_pk=cycle_pk)


@login_required
def report_distribute_all(request, cycle_pk):
    if request.method != 'POST':
        return redirect('reports_cycle', cycle_pk=cycle_pk)
    from accounts.models import VerifierAssignment
    profile = request.user.userprofile
    if profile.role not in ('admin', 'verifier'):
        raise PermissionDenied

    cycle = get_object_or_404(AwardCycle, pk=cycle_pk)

    if cycle.is_open:
        messages.error(request, 'Cannot distribute reports while the cycle is still open.')
        return redirect('reports_cycle', cycle_pk=cycle_pk)

    if profile.role == 'admin':
        queryset = cycle.questionnaires.filter(is_submitted=True, is_distributed=False)
    else:
        assigned_org_ids = VerifierAssignment.objects.filter(
            verifier=request.user
        ).values_list('organization_id', flat=True)
        queryset = cycle.questionnaires.filter(
            is_submitted=True, is_distributed=False,
            organization_id__in=assigned_org_ids,
        )

    now = timezone.now()
    count = 0
    # Organizations configured with a judging stage cannot be released until
    # a judge has completed that stage. Normal organizations remain unchanged.
    queryset = queryset.filter(
        Q(organization__requires_judging=False) |
        Q(organization__requires_judging=True, judging_completed=True)
    )

    for q in queryset.select_related('organization'):
        q.is_distributed = True
        q.distributed_at = now
        q.distributed_by = request.user
        q.save(update_fields=['is_distributed', 'distributed_at', 'distributed_by'])
        AuditLog.log(
            request, 'report.distributed',
            f'Report distributed for {q.organization.name} ({cycle.name})',
            object_repr=str(q.organization),
            object_type='Questionnaire',
            object_id=q.pk,
            organization=q.organization,
        )
        count += 1

    messages.success(request, f'{count} report(s) distributed.')
    return redirect('reports_cycle', cycle_pk=cycle_pk)


@login_required
def report_revoke(request, cycle_pk, q_pk):
    if request.method != 'POST':
        return redirect('reports_cycle', cycle_pk=cycle_pk)
    profile = request.user.userprofile
    if profile.role != 'admin':
        raise PermissionDenied

    cycle = get_object_or_404(AwardCycle, pk=cycle_pk)
    questionnaire = get_object_or_404(Questionnaire, pk=q_pk, cycle=cycle)

    questionnaire.is_distributed = False
    questionnaire.distributed_at = None
    questionnaire.distributed_by = None
    questionnaire.save(update_fields=['is_distributed', 'distributed_at', 'distributed_by'])

    AuditLog.log(
        request, 'report.revoked',
        f'Report access revoked for {questionnaire.organization.name} ({cycle.name})',
        object_repr=str(questionnaire.organization),
        object_type='Questionnaire',
        object_id=questionnaire.pk,
        organization=questionnaire.organization,
    )
    messages.success(request, f'Report access revoked for {questionnaire.organization.name}.')
    return redirect('reports_cycle', cycle_pk=cycle_pk)
