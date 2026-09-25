"""Tests for the report analytics engine.

Kept out of the 3k-line tests.py because this is the maths that decides the
award and it deserves to be findable. Covers the three ranking decisions taken
deliberately — flat mean over criteria, unanswered excluded but surfaced as
completeness, informal sector ranked as its own league — plus the evidence
layer that lets a report say WHY a score was given.
"""
from django.test import TestCase
from django.contrib.auth.models import User
from django.urls import reverse
from django.utils import timezone

from accounts.models import Organization
from .analytics import (
    compute_cycle_analytics, compute_org_analytics, compute_org_benchmark,
)
from .models import (
    AssessmentCategory, AwardCycle, Criterion, LevelIndicator, Questionnaire,
    Response, VerifierResponse,
)


class ReportAnalyticsTest(TestCase):
    def setUp(self):
        self.big_org = Organization.objects.create(
            name='Big Org', org_type='ict')
        self.small_org = Organization.objects.create(
            name='Small Org', org_type='manufacturing')
        self.informal_org = Organization.objects.create(
            name='Informal Org', org_type='informal_sector')

        self.cycle = AwardCycle.objects.create(
            year=2070, name='DSE 2070', is_open=True)

        # Governance holds 3 criteria, Reporting holds 1. Under a mean of
        # category means those two carry equal weight; under a flat mean over
        # criteria Governance carries three times the weight.
        self.gov = AssessmentCategory.objects.create(
            cycle=self.cycle, name='Governance', order=1, is_active=True)
        self.rep = AssessmentCategory.objects.create(
            cycle=self.cycle, name='Reporting', order=2, is_active=True)
        self.informal_cat = AssessmentCategory.objects.create(
            cycle=self.cycle, name='Informal Sector', order=3,
            is_active=True, is_informal_sector_only=True)

        self.g1 = Criterion.objects.create(
            category=self.gov, number=1, name='G1', is_active=True)
        self.g2 = Criterion.objects.create(
            category=self.gov, number=2, name='G2', is_active=True)
        self.g3 = Criterion.objects.create(
            category=self.gov, number=3, name='G3', is_active=True)
        self.r1 = Criterion.objects.create(
            category=self.rep, number=1, name='R1', is_active=True)
        self.i1 = Criterion.objects.create(
            category=self.informal_cat, number=1, name='I1', is_active=True)

        for level, text in enumerate([
            'Nothing in place',
            'Ad hoc',
            'Policy drafted',
            'Policy approved and communicated',
            'Policy monitored',
            'Policy independently audited',
        ]):
            LevelIndicator.objects.create(
                criterion=self.g1, level=level, indicator=text)

        self.q_big = Questionnaire.objects.get(
            cycle=self.cycle, organization=self.big_org)
        self.q_small = Questionnaire.objects.get(
            cycle=self.cycle, organization=self.small_org)
        self.q_informal = Questionnaire.objects.get(
            cycle=self.cycle, organization=self.informal_org)
        now = timezone.now()
        for q in (self.q_big, self.q_small, self.q_informal):
            q.is_submitted = True
            q.submitted_at = now
            q.save()

        self.v1 = User.objects.create_user(
            username='rep_v1', password='pass',
            first_name='Asha', last_name='Mwakalinga')
        self.v2 = User.objects.create_user(
            username='rep_v2', password='pass',
            first_name='Juma', last_name='Kileo')

        # Big Org: strong across the 3 Governance criteria, zero on Reporting.
        for crit in (self.g1, self.g2, self.g3):
            VerifierResponse.objects.create(
                questionnaire=self.q_big, verifier=self.v1,
                criterion=crit, score=4)
        VerifierResponse.objects.create(
            questionnaire=self.q_big, verifier=self.v1,
            criterion=self.r1, score=0)

    def _entry(self, result, criterion):
        for cat_entry in result['categories']:
            for entry in cat_entry['criteria']:
                if entry['criterion'].pk == criterion.pk:
                    return entry
        raise AssertionError(f'criterion {criterion} missing from analytics')

    def _league(self, result, key):
        return next(l for l in result['leagues'] if l['key'] == key)

    # ── Flat mean over criteria ──────────────────────────────────────────

    def test_overall_is_flat_mean_over_criteria(self):
        result = compute_org_analytics(self.q_big)
        # Flat over 4 criteria: (4+4+4+0)/4 = 3.0
        # A mean of category means would give (4.0 + 0.0)/2 = 2.0
        self.assertEqual(result['overall_verified_score'], 3.0)

    def test_small_category_cannot_outweigh_large_one(self):
        result = compute_org_analytics(self.q_big)
        self.assertGreater(result['overall_verified_score'], 2.5)

    # ── Completeness ─────────────────────────────────────────────────────

    def test_unanswered_criteria_excluded_from_score(self):
        Response.objects.create(
            questionnaire=self.q_big, criterion=self.g1, score=5)
        result = compute_org_analytics(self.q_big)
        self.assertEqual(result['overall_claimed_score'], 5.0)

    def test_completeness_reports_what_the_score_hides(self):
        Response.objects.create(
            questionnaire=self.q_big, criterion=self.g1, score=5)
        result = compute_org_analytics(self.q_big)
        self.assertEqual(result['answered_count'], 1)
        self.assertEqual(result['criteria_count'], 4)
        self.assertEqual(result['completeness'], 25.0)

    def test_completeness_appears_on_every_ranking_row(self):
        Response.objects.create(
            questionnaire=self.q_big, criterion=self.g1, score=5)
        result = compute_cycle_analytics(self.cycle)
        row = next(r for r in result['ranked_orgs'] if r['org'] == self.big_org)
        self.assertEqual(row['completeness'], 25.0)
        self.assertEqual(row['answered'], 1)
        self.assertEqual(row['total'], 4)

    # ── Informal sector league ───────────────────────────────────────────

    def test_informal_org_ranked_in_its_own_league(self):
        VerifierResponse.objects.create(
            questionnaire=self.q_informal, verifier=self.v1,
            criterion=self.i1, score=5)
        result = compute_cycle_analytics(self.cycle)
        self.assertEqual(
            {l['key'] for l in result['leagues']}, {'main', 'informal'})

        informal = self._league(result, 'informal')
        self.assertEqual(
            [r['org'] for r in informal['ranked_orgs']], [self.informal_org])
        self.assertEqual(informal['ranked_orgs'][0]['rank'], 1)

        main = self._league(result, 'main')
        self.assertNotIn(
            self.informal_org, [r['org'] for r in main['ranked_orgs']])

    def test_informal_league_scored_only_on_its_own_categories(self):
        VerifierResponse.objects.create(
            questionnaire=self.q_informal, verifier=self.v1,
            criterion=self.i1, score=5)
        result = compute_cycle_analytics(self.cycle)
        row = self._league(result, 'informal')['ranked_orgs'][0]
        self.assertEqual(row['total'], 1)
        self.assertEqual(row['overall_verified_score'], 5.0)

    # ── The "because" ────────────────────────────────────────────────────

    def test_verifier_notes_reach_the_report(self):
        VerifierResponse.objects.filter(
            questionnaire=self.q_big, criterion=self.g2
        ).update(notes='Draft sighted, no board minutes.')
        entry = self._entry(compute_org_analytics(self.q_big), self.g2)
        self.assertEqual(len(entry['verifier_notes']), 1)
        self.assertEqual(
            entry['verifier_notes'][0]['notes'],
            'Draft sighted, no board minutes.')
        self.assertEqual(entry['verifier_notes'][0]['name'], 'Asha Mwakalinga')

    def test_each_verifier_named_with_their_own_score(self):
        VerifierResponse.objects.create(
            questionnaire=self.q_big, verifier=self.v2,
            criterion=self.g1, score=2, notes='Weaker than claimed.')
        entry = self._entry(compute_org_analytics(self.q_big), self.g1)
        by_name = {v['name']: v['score'] for v in entry['verifier_entries']}
        self.assertEqual(by_name['Asha Mwakalinga'], 4)
        self.assertEqual(by_name['Juma Kileo'], 2)
        self.assertEqual(entry['verified_score'], 3.0)

    def test_level_indicator_explains_score_and_next_step(self):
        entry = self._entry(compute_org_analytics(self.q_big), self.g1)
        self.assertEqual(entry['level_now'], 4)
        self.assertEqual(entry['level_now_text'], 'Policy monitored')
        self.assertEqual(entry['level_next'], 5)
        self.assertEqual(
            entry['level_next_text'], 'Policy independently audited')

    def test_level_floors_rather_than_rounds_up(self):
        VerifierResponse.objects.create(
            questionnaire=self.q_big, verifier=self.v2,
            criterion=self.g1, score=3)
        entry = self._entry(compute_org_analytics(self.q_big), self.g1)
        self.assertEqual(entry['verified_score'], 3.5)
        self.assertEqual(entry['level_now'], 3)
        self.assertEqual(
            entry['level_now_text'], 'Policy approved and communicated')

    def test_top_level_has_no_next_step(self):
        VerifierResponse.objects.filter(
            questionnaire=self.q_big, criterion=self.g1).update(score=5)
        entry = self._entry(compute_org_analytics(self.q_big), self.g1)
        self.assertEqual(entry['level_now'], 5)
        self.assertIsNone(entry['level_next'])
        self.assertIsNone(entry['level_next_text'])

    # ── Verifier disagreement ────────────────────────────────────────────

    def test_verifier_disagreement_flagged(self):
        VerifierResponse.objects.create(
            questionnaire=self.q_big, verifier=self.v2,
            criterion=self.g1, score=1)
        result = compute_org_analytics(self.q_big)
        entry = self._entry(result, self.g1)
        self.assertEqual(entry['spread'], 3)
        self.assertTrue(entry['is_disputed'])
        self.assertIn(entry, result['disputed_criteria'])

    def test_close_verifier_scores_not_flagged(self):
        VerifierResponse.objects.create(
            questionnaire=self.q_big, verifier=self.v2,
            criterion=self.g1, score=3)
        entry = self._entry(compute_org_analytics(self.q_big), self.g1)
        self.assertEqual(entry['spread'], 1)
        self.assertFalse(entry['is_disputed'])

    # ── Gap in words ─────────────────────────────────────────────────────

    def test_over_stated_gap_described_in_words(self):
        for crit in (self.g1, self.g2, self.g3, self.r1):
            Response.objects.create(
                questionnaire=self.q_big, criterion=crit, score=5)
        result = compute_org_analytics(self.q_big)
        self.assertEqual(result['gap_direction'], 'over')
        self.assertIn('higher than verified', result['gap_phrase'])

    def test_matching_self_assessment_described_as_matched(self):
        for crit, score in (
            (self.g1, 4), (self.g2, 4), (self.g3, 4), (self.r1, 0),
        ):
            Response.objects.create(
                questionnaire=self.q_big, criterion=crit, score=score)
        result = compute_org_analytics(self.q_big)
        self.assertEqual(result['gap_direction'], 'match')

    def test_under_stated_gap_described_in_words(self):
        for crit in (self.g1, self.g2, self.g3, self.r1):
            Response.objects.create(
                questionnaire=self.q_big, criterion=crit, score=0)
        result = compute_org_analytics(self.q_big)
        self.assertEqual(result['gap_direction'], 'under')
        self.assertIn('lower than verified', result['gap_phrase'])

    # ── Priorities and strengths ─────────────────────────────────────────

    def test_priorities_are_the_weakest_criteria(self):
        result = compute_org_analytics(self.q_big)
        self.assertEqual(result['priority_criteria'][0]['criterion'], self.r1)

    def test_strengths_are_criteria_at_level_four_or_above(self):
        result = compute_org_analytics(self.q_big)
        names = {e['criterion'].name for e in result['strength_criteria']}
        self.assertEqual(names, {'G1', 'G2', 'G3'})

    def test_strong_org_still_gets_a_next_step(self):
        VerifierResponse.objects.filter(questionnaire=self.q_big).update(score=5)
        result = compute_org_analytics(self.q_big)
        self.assertTrue(result['priority_criteria'])

    # ── Per-category ranking ─────────────────────────────────────────────

    def test_every_org_ranked_within_every_category(self):
        for crit in (self.g1, self.g2, self.g3):
            VerifierResponse.objects.create(
                questionnaire=self.q_small, verifier=self.v1,
                criterion=crit, score=1)
        VerifierResponse.objects.create(
            questionnaire=self.q_small, verifier=self.v1,
            criterion=self.r1, score=5)

        result = compute_cycle_analytics(self.cycle)
        by_cat = {e['category'].name: e for e in result['category_table']}

        gov = by_cat['Governance']['placings']
        self.assertEqual(gov[0]['org'], self.big_org)
        self.assertEqual(gov[0]['rank'], 1)
        self.assertEqual(gov[1]['org'], self.small_org)
        self.assertEqual(gov[1]['rank'], 2)

        # Reporting inverts the order — the whole point of per-category rank.
        rep = by_cat['Reporting']['placings']
        self.assertEqual(rep[0]['org'], self.small_org)
        self.assertEqual(rep[1]['org'], self.big_org)

    def test_equal_scores_share_a_rank(self):
        for crit, score in (
            (self.g1, 4), (self.g2, 4), (self.g3, 4), (self.r1, 0),
        ):
            VerifierResponse.objects.create(
                questionnaire=self.q_small, verifier=self.v1,
                criterion=crit, score=score)
        result = compute_cycle_analytics(self.cycle)
        main = self._league(result, 'main')
        self.assertEqual([r['rank'] for r in main['ranked_orgs']], [1, 1])

    def test_unscored_org_has_no_rank_rather_than_last_place(self):
        result = compute_cycle_analytics(self.cycle)
        main = self._league(result, 'main')
        small = next(
            r for r in main['ranked_orgs'] if r['org'] == self.small_org)
        self.assertIsNone(small['rank'])
        self.assertIsNone(small['overall_verified_score'])
        self.assertEqual(main['ranked_orgs'][0]['org'], self.big_org)

    # ── Benchmark leaks no peer data ─────────────────────────────────────

    def test_benchmark_exposes_position_and_averages_only(self):
        for crit in (self.g1, self.g2, self.g3, self.r1):
            VerifierResponse.objects.create(
                questionnaire=self.q_small, verifier=self.v1,
                criterion=crit, score=2)

        benchmark = compute_org_benchmark(self.q_big)
        self.assertEqual(benchmark['rank'], 1)
        self.assertEqual(benchmark['league_size'], 2)
        self.assertEqual(benchmark['cohort_avg'], 2.5)  # avg(3.0, 2.0)
        # Members read the org report — no peer may be named in it.
        self.assertNotIn('Small Org', repr(benchmark))

    def test_benchmark_gives_org_report_its_category_positions(self):
        for crit in (self.g1, self.g2, self.g3):
            VerifierResponse.objects.create(
                questionnaire=self.q_small, verifier=self.v1,
                criterion=crit, score=1)
        VerifierResponse.objects.create(
            questionnaire=self.q_small, verifier=self.v1,
            criterion=self.r1, score=5)

        benchmark = compute_org_benchmark(self.q_big)
        result = compute_org_analytics(self.q_big, benchmark=benchmark)
        by_name = {e['category'].name: e for e in result['categories']}
        self.assertEqual(by_name['Governance']['rank'], 1)
        self.assertEqual(by_name['Reporting']['rank'], 2)
        self.assertEqual(by_name['Governance']['rank_of'], 2)

    def test_criterion_carries_cohort_comparison(self):
        for crit in (self.g1, self.g2, self.g3, self.r1):
            VerifierResponse.objects.create(
                questionnaire=self.q_small, verifier=self.v1,
                criterion=crit, score=2)
        benchmark = compute_org_benchmark(self.q_big)
        result = compute_org_analytics(self.q_big, benchmark=benchmark)
        entry = self._entry(result, self.g1)
        self.assertEqual(entry['cohort_avg'], 3.0)  # avg(4, 2)
        self.assertEqual(entry['vs_cohort'], 1.0)

    # ── Numeric criteria stay out of the score ───────────────────────────

    def test_numeric_criteria_excluded_from_averages(self):
        numeric = Criterion.objects.create(
            category=self.gov, number=9, name='Headcount', is_active=True,
            is_numeric=True,
            numeric_fields=[{'name': 'Staff', 'type': 'integer'}])
        Response.objects.create(
            questionnaire=self.q_big, criterion=numeric,
            numeric_data={'Staff': 40})
        result = compute_org_analytics(self.q_big)
        self.assertEqual(result['overall_verified_score'], 3.0)
        # It still counts toward how complete the submission is.
        self.assertEqual(result['criteria_count'], 5)
        self.assertEqual(result['answered_count'], 1)


class ReportPageRenderTest(TestCase):
    """The pages themselves â€” the previous criterion detail was rendered but
    permanently hidden by CSS, so 'it renders' is not enough on its own."""

    def setUp(self):
        from accounts.models import UserProfile, VerifierAssignment

        self.org = Organization.objects.create(name='Render Org', org_type='ict')
        self.peer = Organization.objects.create(name='Peer Org', org_type='ict')
        self.cycle = AwardCycle.objects.create(
            year=2071, name='DSE 2071', is_open=True)

        self.cat = AssessmentCategory.objects.create(
            cycle=self.cycle, name='Governance', order=1, is_active=True)
        self.c1 = Criterion.objects.create(
            category=self.cat, number=1, name='Board oversight', is_active=True)
        self.c2 = Criterion.objects.create(
            category=self.cat, number=2, name='Staff policy', is_active=True)
        for level, text in enumerate([
            'Nothing in place', 'Ad hoc', 'Policy drafted',
            'Policy approved', 'Policy monitored', 'Policy audited',
        ]):
            LevelIndicator.objects.create(
                criterion=self.c2, level=level, indicator=text)

        self.q = Questionnaire.objects.get(
            cycle=self.cycle, organization=self.org)
        self.q_peer = Questionnaire.objects.get(
            cycle=self.cycle, organization=self.peer)
        now = timezone.now()
        for q in (self.q, self.q_peer):
            q.is_submitted = True
            q.submitted_at = now
            q.save()

        self.verifier = User.objects.create_user(
            username='render_v', password='Pass123!',
            first_name='Asha', last_name='Mwakalinga')
        UserProfile.objects.create(user=self.verifier, role='verifier')
        VerifierAssignment.objects.create(
            verifier=self.verifier, organization=self.org)

        self.admin = User.objects.create_user(
            username='render_admin', password='Pass123!')
        UserProfile.objects.create(user=self.admin, role='admin')

        self.member = User.objects.create_user(
            username='render_member', password='Pass123!')
        UserProfile.objects.create(
            user=self.member, role='member', organization=self.org)

        Response.objects.create(
            questionnaire=self.q, criterion=self.c1, score=5,
            notes='We hold quarterly board meetings.')
        Response.objects.create(
            questionnaire=self.q, criterion=self.c2, score=5)
        VerifierResponse.objects.create(
            questionnaire=self.q, verifier=self.verifier, criterion=self.c1,
            score=4, notes='Minutes sighted for three of four quarters.')
        VerifierResponse.objects.create(
            questionnaire=self.q, verifier=self.verifier, criterion=self.c2,
            score=2, notes='Draft sighted, no board approval evidenced.')
        for crit in (self.c1, self.c2):
            VerifierResponse.objects.create(
                questionnaire=self.q_peer, verifier=self.verifier,
                criterion=crit, score=1)

    def _get_report(self, username='render_admin'):
        self.client.login(username=username, password='Pass123!')
        return self.client.get(
            reverse('org_report', args=[self.cycle.pk, self.q.pk]))

    def test_org_report_renders(self):
        self.assertEqual(self._get_report().status_code, 200)

    def test_report_leads_with_verified_score(self):
        html = self._get_report().content.decode()
        self.assertIn('Verified Score', html)
        self.assertIn('Self-assessed', html)

    def test_report_states_the_gap_in_words(self):
        html = self._get_report().content.decode()
        # Claimed 5 and 5, verified 4 and 2 â€” a clear over-statement.
        self.assertIn('higher than verified', html)

    def test_report_shows_the_verifier_reason(self):
        html = self._get_report().content.decode()
        self.assertIn('Draft sighted, no board approval evidenced.', html)
        self.assertIn('Asha Mwakalinga', html)

    def test_report_shows_why_this_level_and_the_next_step(self):
        html = self._get_report().content.decode()
        self.assertIn('At level 2 because', html)
        self.assertIn('Policy drafted', html)
        self.assertIn('To reach level 3', html)
        self.assertIn('Policy approved', html)

    def test_criterion_detail_uses_a_class_the_stylesheet_defines(self):
        """Regression: the old markup toggled .cat-detail.open, a rule that
        does not exist, so this section could never be opened."""
        html = self._get_report().content.decode()
        self.assertIn('crit-group-body', html)
        self.assertNotIn("classList.toggle('open')", html)

    def test_report_shows_rank_and_cohort(self):
        html = self._get_report().content.decode()
        self.assertIn('Ranked', html)
        self.assertIn('Cohort average', html)

    def test_member_sees_own_report_without_peer_names(self):
        self.q.is_distributed = True
        self.q.save()
        self.cycle.is_open = False
        self.cycle.save()
        self.client.login(username='render_member', password='Pass123!')
        response = self.client.get(
            reverse('org_report', args=[self.cycle.pk, self.q.pk]))
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()
        self.assertIn('Ranked', html)
        # Rank and cohort averages are safe; a peer's identity is not.
        self.assertNotIn('Peer Org', html)

    def test_member_does_not_see_verifier_disagreements_panel(self):
        VerifierResponse.objects.create(
            questionnaire=self.q,
            verifier=User.objects.create_user(username='render_v2', password='x'),
            criterion=self.c1, score=0)
        self.q.is_distributed = True
        self.q.save()
        self.cycle.is_open = False
        self.cycle.save()
        self.client.login(username='render_member', password='Pass123!')
        html = self.client.get(
            reverse('org_report', args=[self.cycle.pk, self.q.pk])
        ).content.decode()
        self.assertNotIn('Verifier Disagreements', html)

        self.client.login(username='render_admin', password='Pass123!')
        html = self.client.get(
            reverse('org_report', args=[self.cycle.pk, self.q.pk])
        ).content.decode()
        self.assertIn('Verifier Disagreements', html)

    def test_cycle_results_renders_leagues_and_category_standings(self):
        self.client.login(username='render_admin', password='Pass123!')
        response = self.client.get(
            reverse('cycle_results', args=[self.cycle.pk]))
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()
        self.assertIn('Main League', html)
        self.assertIn('Standings by Category', html)
        self.assertIn('Completed', html)

    def test_excel_export_carries_verifier_reasoning(self):
        import io
        import openpyxl

        self.client.login(username='render_admin', password='Pass123!')
        response = self.client.get(
            reverse('export_org_excel', args=[self.cycle.pk, self.q.pk]))
        self.assertEqual(response.status_code, 200)

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        text = '\n'.join(
            str(cell.value)
            for row in wb.active.iter_rows()
            for cell in row
            if cell.value is not None
        )
        self.assertIn('Draft sighted, no board approval evidenced.', text)
        self.assertIn('Policy approved', text)
        self.assertIn('Verified Score (out of 5)', text)

    def test_cycle_excel_export_has_a_by_category_sheet(self):
        import io
        import openpyxl

        self.client.login(username='render_admin', password='Pass123!')
        response = self.client.get(
            reverse('export_cycle_excel', args=[self.cycle.pk]))
        self.assertEqual(response.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertIn('By Category', wb.sheetnames)
        self.assertIn('Rankings', wb.sheetnames)

    def test_pdf_export_still_builds(self):
        self.client.login(username='render_admin', password='Pass123!')
        response = self.client.get(
            reverse('export_org_pdf', args=[self.cycle.pk, self.q.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF'))



class CompareScoringCommandTest(TestCase):
    """The command exists to be run against production, so the thing that
    matters most is that it writes nothing."""

    def setUp(self):
        self.org = Organization.objects.create(name='Cmp Org', org_type='ict')
        self.cycle = AwardCycle.objects.create(
            year=2075, name='DSE 2075', is_open=True)
        self.cat = AssessmentCategory.objects.create(
            cycle=self.cycle, name='Cat', order=1, is_active=True)
        self.c1 = Criterion.objects.create(
            category=self.cat, number=1, name='C1', is_active=True)
        self.q = Questionnaire.objects.get(
            cycle=self.cycle, organization=self.org)
        self.q.is_submitted = True
        self.q.submitted_at = timezone.now()
        self.q.save()
        self.v = User.objects.create_user(username='cmp_v', password='pass')
        Response.objects.create(
            questionnaire=self.q, criterion=self.c1, score=5)
        VerifierResponse.objects.create(
            questionnaire=self.q, verifier=self.v, criterion=self.c1, score=3)

    def _run(self, **kwargs):
        from io import StringIO
        from django.core.management import call_command
        out = StringIO()
        call_command('compare_scoring', stdout=out, **kwargs)
        return out.getvalue()

    def test_runs_and_reports_the_cycle(self):
        output = self._run()
        self.assertIn('DSE 2075', output)
        self.assertIn('Cmp Org', output)
        self.assertIn('No data was modified', output)

    def test_writes_nothing(self):
        before = {
            'responses': list(Response.objects.values_list('id', 'score')),
            'verifier': list(VerifierResponse.objects.values_list('id', 'score')),
            'questionnaires': list(
                Questionnaire.objects.values_list('id', 'is_submitted', 'is_distributed')),
            'cycles': list(AwardCycle.objects.values_list('id', 'is_open')),
        }
        self._run()
        self.assertEqual(before['responses'],
                         list(Response.objects.values_list('id', 'score')))
        self.assertEqual(before['verifier'],
                         list(VerifierResponse.objects.values_list('id', 'score')))
        self.assertEqual(before['questionnaires'], list(
            Questionnaire.objects.values_list('id', 'is_submitted', 'is_distributed')))
        self.assertEqual(before['cycles'],
                         list(AwardCycle.objects.values_list('id', 'is_open')))

    def test_cycle_filter_excludes_others(self):
        other_org = Organization.objects.create(name='Other Cmp Org', org_type='ict')
        AwardCycle.objects.create(year=2076, name='DSE 2076', is_open=False)
        output = self._run(cycle=2075)
        self.assertIn('DSE 2075', output)
        self.assertNotIn('DSE 2076', output)

    def test_distributed_only_filter(self):
        output = self._run(distributed_only=True)
        # Nothing is distributed, so the cycle contributes no rows.
        self.assertIn('questionnaires examined : 0', output)

