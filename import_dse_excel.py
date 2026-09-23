import openpyxl
from django.db import transaction
from assessment.models import (
    AwardCycle,
    AssessmentCategory,
    Criterion,
    LevelIndicator,
)

EXCEL_FILE = "/app/DSE_upload_ready.xlsx"
YEAR = 2026
CYCLE_NAME = "2026 DSE Members Award"

print("=" * 70)
print("DSE EXCEL IMPORT")
print("=" * 70)

# ---------------------------------------------------------
# Find or create the Award Cycle
# ---------------------------------------------------------
cycle, created = AwardCycle.objects.get_or_create(
    year=YEAR,
    defaults={
        "name": CYCLE_NAME,
        "is_open": False,
    }
)

if created:
    print(f"Created Award Cycle: {cycle.name}")
else:
    print(f"Using existing Award Cycle: {cycle.name}")

# ---------------------------------------------------------
# Read Excel
# ---------------------------------------------------------
wb = openpyxl.load_workbook(
    EXCEL_FILE,
    data_only=True
)

print(f"Excel sheets found: {wb.sheetnames}")
print()

total_categories = 0
total_criteria = 0
total_levels = 0

# ---------------------------------------------------------
# Process each organization-type sheet
# ---------------------------------------------------------
for sheet_name in wb.sheetnames:

    # Ignore the summary/detail sheets
    if sheet_name in [
        "Assessment",
        "Source Details",
    ]:
        continue

    ws = wb[sheet_name]

    print("-" * 70)
    print(f"PROCESSING SHEET: {sheet_name}")
    print("-" * 70)

    # -----------------------------------------------------
    # Find the header row
    # -----------------------------------------------------
    header_row = None

    for row_number in range(1, min(ws.max_row, 20) + 1):
        values = [
            str(ws.cell(row=row_number, column=col).value or "").strip()
            for col in range(1, min(ws.max_column, 13) + 1)
        ]

        if (
            "Assessment Category" in values
            and "Assessment Criteria" in values
            and "Measurable Parameters" in values
        ):
            header_row = row_number
            break

    if not header_row:
        print(f"WARNING: No valid header found in {sheet_name}")
        continue

    # -----------------------------------------------------
    # Map columns
    # -----------------------------------------------------
    headers = {}

    for col in range(1, ws.max_column + 1):
        value = ws.cell(header_row, col).value

        if value:
            headers[str(value).strip()] = col

    required = [
        "Assessment Category",
        "Assessment Areas",
        "Assessment Criteria",
        "Measurable Parameters",
    ]

    missing = [x for x in required if x not in headers]

    if missing:
        print(f"WARNING: Missing columns: {missing}")
        continue

    category_col = headers["Assessment Category"]
    area_col = headers["Assessment Areas"]
    criterion_col = headers["Assessment Criteria"]
    parameter_col = headers["Measurable Parameters"]

    regulation_col = headers.get(
        "Relevant CMSA Act or Guideline /DSE Rule/Applicable Regulation"
    )

    weight_col = headers.get("Weight (%)")

    # -----------------------------------------------------
    # Track current values because Excel uses blank cells
    # for repeated values.
    # -----------------------------------------------------
    current_category = None
    current_area = None
    current_criterion = None

    category_order = {}
    criterion_order = {}

    # -----------------------------------------------------
    # Process rows
    # -----------------------------------------------------
    for row in range(header_row + 1, ws.max_row + 1):

        category_value = ws.cell(row, category_col).value
        area_value = ws.cell(row, area_col).value
        criterion_value = ws.cell(row, criterion_col).value
        parameter_value = ws.cell(row, parameter_col).value

        if category_value is not None:
            current_category = str(category_value).strip()

        if area_value is not None:
            current_area = str(area_value).strip()

        if criterion_value is not None:
            current_criterion = str(criterion_value).strip()

        if not current_category:
            continue

        if not current_criterion:
            continue

        # -------------------------------------------------
        # Category
        #
        # Include the Excel sheet name in the category so
        # Bond Issuer, Bond Trader, Custodian, etc. remain
        # separate without changing the database.
        # -------------------------------------------------
        category_name = f"{sheet_name} - {current_category}"

        if category_name not in category_order:
            category_order[category_name] = len(category_order)

        category, category_created = (
            AssessmentCategory.objects.get_or_create(
                cycle=cycle,
                name=category_name,
                defaults={
                    "description": (
                        f"Imported from Excel sheet: {sheet_name}. "
                        f"Assessment Category: {current_category}"
                    ),
                    "order": category_order[category_name],
                    "is_active": True,
                },
            )
        )

        if category_created:
            total_categories += 1

        # -------------------------------------------------
        # Criterion number
        # -------------------------------------------------
        criterion_key = (
            category.id,
            current_criterion,
        )

        if criterion_key not in criterion_order:
            criterion_order[criterion_key] = (
                len([
                    x for x in criterion_order
                    if x[0] == category.id
                ]) + 1
            )

        criterion_number = criterion_order[criterion_key]

        # -------------------------------------------------
        # Preserve information that does not have a
        # dedicated database field.
        # -------------------------------------------------
        description_parts = []

        if current_area:
            description_parts.append(
                f"Assessment Area: {current_area}"
            )

        if regulation_col:
            regulation = ws.cell(row, regulation_col).value

            if regulation:
                description_parts.append(
                    f"Applicable Regulation: {str(regulation).strip()}"
                )

        if weight_col:
            weight = ws.cell(row, weight_col).value

            if weight is not None:
                description_parts.append(
                    f"Weight (%): {weight}"
                )

        description = "\n".join(description_parts)

        # -------------------------------------------------
        # Criterion
        # -------------------------------------------------
        criterion, criterion_created = (
            Criterion.objects.update_or_create(
                category=category,
                number=criterion_number,
                defaults={
                    "name": current_criterion,
                    "description": description,
                    "order": criterion_number,
                    "is_active": True,
                    "is_numeric": False,
                    "numeric_fields": [],
                },
            )
        )

        if criterion_created:
            total_criteria += 1

        # -------------------------------------------------
        # Measurable Parameter → Level Indicator
        #
        # Existing LevelIndicator only accepts levels 0-5.
        # Therefore parameters under the same criterion are
        # assigned levels 1, 2, 3... without changing the DB.
        # -------------------------------------------------
        if parameter_value is not None:

            parameter = str(parameter_value).strip()

            if not parameter:
                continue

            existing_count = (
                LevelIndicator.objects
                .filter(criterion=criterion)
                .count()
            )

            level = existing_count + 1

docker exec -i dse_django sh -c 'cat > /app/import_dse_excel.py' <<'PY'
import openpyxl
from django.db import transaction
from assessment.models import (
    AwardCycle,
    AssessmentCategory,
    Criterion,
    LevelIndicator,
)

EXCEL_FILE = "/app/DSE_upload_ready.xlsx"
YEAR = 2026
CYCLE_NAME = "2026 DSE Members Award"

print("=" * 70)
print("DSE EXCEL IMPORT")
print("=" * 70)

# ---------------------------------------------------------
# Find or create the Award Cycle
# ---------------------------------------------------------
cycle, created = AwardCycle.objects.get_or_create(
    year=YEAR,
    defaults={
        "name": CYCLE_NAME,
        "is_open": False,
    }
)

if created:
    print(f"Created Award Cycle: {cycle.name}")
else:
    print(f"Using existing Award Cycle: {cycle.name}")

# ---------------------------------------------------------
# Read Excel
# ---------------------------------------------------------
wb = openpyxl.load_workbook(
    EXCEL_FILE,
    data_only=True
)

print(f"Excel sheets found: {wb.sheetnames}")
print()

total_categories = 0
total_criteria = 0
total_levels = 0

# ---------------------------------------------------------
# Process each organization-type sheet
# ---------------------------------------------------------
for sheet_name in wb.sheetnames:

    # Ignore the summary/detail sheets
    if sheet_name in [
        "Assessment",
        "Source Details",
    ]:
        continue

    ws = wb[sheet_name]

    print("-" * 70)
    print(f"PROCESSING SHEET: {sheet_name}")
    print("-" * 70)

    # -----------------------------------------------------
    # Find the header row
    # -----------------------------------------------------
    header_row = None

    for row_number in range(1, min(ws.max_row, 20) + 1):
        values = [
            str(ws.cell(row=row_number, column=col).value or "").strip()
            for col in range(1, min(ws.max_column, 13) + 1)
        ]

        if (
            "Assessment Category" in values
            and "Assessment Criteria" in values
            and "Measurable Parameters" in values
        ):
            header_row = row_number
            break

    if not header_row:
        print(f"WARNING: No valid header found in {sheet_name}")
        continue

    # -----------------------------------------------------
    # Map columns
    # -----------------------------------------------------
    headers = {}

    for col in range(1, ws.max_column + 1):
        value = ws.cell(header_row, col).value

        if value:
            headers[str(value).strip()] = col

    required = [
        "Assessment Category",
        "Assessment Areas",
        "Assessment Criteria",
        "Measurable Parameters",
    ]

    missing = [x for x in required if x not in headers]

    if missing:
        print(f"WARNING: Missing columns: {missing}")
        continue

    category_col = headers["Assessment Category"]
    area_col = headers["Assessment Areas"]
    criterion_col = headers["Assessment Criteria"]
    parameter_col = headers["Measurable Parameters"]

    regulation_col = headers.get(
        "Relevant CMSA Act or Guideline /DSE Rule/Applicable Regulation"
    )

    weight_col = headers.get("Weight (%)")

    # -----------------------------------------------------
    # Track current values because Excel uses blank cells
    # for repeated values.
    # -----------------------------------------------------
    current_category = None
    current_area = None
    current_criterion = None

    category_order = {}
    criterion_order = {}

    # -----------------------------------------------------
    # Process rows
    # -----------------------------------------------------
    for row in range(header_row + 1, ws.max_row + 1):

        category_value = ws.cell(row, category_col).value
        area_value = ws.cell(row, area_col).value
        criterion_value = ws.cell(row, criterion_col).value
        parameter_value = ws.cell(row, parameter_col).value

        if category_value is not None:
            current_category = str(category_value).strip()

        if area_value is not None:
            current_area = str(area_value).strip()

        if criterion_value is not None:
            current_criterion = str(criterion_value).strip()

        if not current_category:
            continue

        if not current_criterion:
            continue

        # -------------------------------------------------
        # Category
        #
        # Include the Excel sheet name in the category so
        # Bond Issuer, Bond Trader, Custodian, etc. remain
        # separate without changing the database.
        # -------------------------------------------------
        category_name = f"{sheet_name} - {current_category}"

        if category_name not in category_order:
            category_order[category_name] = len(category_order)

        category, category_created = (
            AssessmentCategory.objects.get_or_create(
                cycle=cycle,
                name=category_name,
                defaults={
                    "description": (
                        f"Imported from Excel sheet: {sheet_name}. "
                        f"Assessment Category: {current_category}"
                    ),
                    "order": category_order[category_name],
                    "is_active": True,
                },
            )
        )

        if category_created:
            total_categories += 1

        # -------------------------------------------------
        # Criterion number
        # -------------------------------------------------
        criterion_key = (
            category.id,
            current_criterion,
        )

        if criterion_key not in criterion_order:
            criterion_order[criterion_key] = (
                len([
                    x for x in criterion_order
                    if x[0] == category.id
                ]) + 1
            )

        criterion_number = criterion_order[criterion_key]

        # -------------------------------------------------
        # Preserve information that does not have a
        # dedicated database field.
        # -------------------------------------------------
        description_parts = []

        if current_area:
            description_parts.append(
                f"Assessment Area: {current_area}"
            )

        if regulation_col:
            regulation = ws.cell(row, regulation_col).value

            if regulation:
                description_parts.append(
                    f"Applicable Regulation: {str(regulation).strip()}"
                )

        if weight_col:
            weight = ws.cell(row, weight_col).value

            if weight is not None:
                description_parts.append(
                    f"Weight (%): {weight}"
                )

        description = "\n".join(description_parts)

        # -------------------------------------------------
        # Criterion
        # -------------------------------------------------
        criterion, criterion_created = (
            Criterion.objects.update_or_create(
                category=category,
                number=criterion_number,
                defaults={
                    "name": current_criterion,
                    "description": description,
                    "order": criterion_number,
                    "is_active": True,
                    "is_numeric": False,
                    "numeric_fields": [],
                },
            )
        )

        if criterion_created:
            total_criteria += 1

        # -------------------------------------------------
        # Measurable Parameter → Level Indicator
        #
        # Existing LevelIndicator only accepts levels 0-5.
        # Therefore parameters under the same criterion are
        # assigned levels 1, 2, 3... without changing the DB.
        # -------------------------------------------------
        if parameter_value is not None:

            parameter = str(parameter_value).strip()

            if not parameter:
                continue

            existing_count = (
                LevelIndicator.objects
                .filter(criterion=criterion)
                .count()
            )

            level = existing_count + 1

            if level > 5:
                print(
                    f"WARNING: More than 5 parameters for "
                    f"{criterion.name}. Extra parameter skipped."
                )
                continue

            indicator, indicator_created = (
                LevelIndicator.objects.update_or_create(
                    criterion=criterion,
                    level=level,
                    defaults={
                        "indicator": parameter
                    },
                )
            )

            if indicator_created:
                total_levels += 1

print()
print("=" * 70)
print("IMPORT COMPLETED")
print("=" * 70)
print(f"Award Cycle : {cycle.name}")
print(f"Categories  : {total_categories}")
print(f"Criteria    : {total_criteria}")
print(f"Indicators  : {total_levels}")
print("=" * 70)
